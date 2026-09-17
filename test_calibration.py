# -*- coding: utf-8 -*-
"""Calibration provenance, transfer, combination, and export contracts."""
from __future__ import annotations

import math
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook

from run_ste_effects_batch import build_layer
from ste_lab.calibration import (
    PROVENANCE_COMBINED_ESTIMATE,
    PROVENANCE_MEASURED,
    PROVENANCE_TRANSFERRED,
    PROVENANCE_TRANSFERRED_EXTRAPOLATED,
    CalibrationConfig,
    acceptance_counts,
    agreement_metrics,
    apply_log_ratio,
    build_final_rows,
    build_validation_tables,
    combine_empirical_and_modelled,
    extrapolation_distance_semitones,
    legacy_scalar_transfer,
    load_config,
    optimal_weight,
    qa_index,
    qa_status,
    shape_normalised_metrics,
    two_log_ratio_transfer,
)
from ste_lab.calibration_export import write_measured_data
from ste_lab.catalog import family_donor_instrument
from ste_lab.session import Project
from ste_lab.transfer import media_of
from ste_lab.zenodo_export import _note_grid, _write_acoustic_table, measured_ceiling_midi


class TestProvenanceAndDistance(unittest.TestCase):
    def test_measured_stays_measured(self):
        cfg = CalibrationConfig()
        rows = build_final_rows(
            instrument="bass_clarinet",
            dynamic="mf",
            empirical={60: 10.0},
            transfer={60: 99.0},
            cfg=cfg,
        )
        self.assertEqual(rows[0]["provenance"], PROVENANCE_MEASURED)
        self.assertEqual(rows[0]["final_value"], 10.0)
        self.assertEqual(rows[0]["empirical_value"], 10.0)

    def test_transferred_inside_span(self):
        cfg = CalibrationConfig()
        rows = build_final_rows(
            instrument="bass_clarinet",
            dynamic="pp",
            empirical={50: 8.0, 60: 9.0},
            transfer={55: 7.0},
            cfg=cfg,
        )
        hit = next(r for r in rows if r["midi"] == 55)
        self.assertEqual(hit["provenance"], PROVENANCE_TRANSFERRED)
        self.assertEqual(hit["extrapolation_distance_semitones"], 0)

    def test_transferred_extrapolated_outside_span(self):
        cfg = CalibrationConfig()
        rows = build_final_rows(
            instrument="bass_clarinet",
            dynamic="pp",
            empirical={50: 8.0, 71: 9.0},
            transfer={77: 7.0},
            cfg=cfg,
        )
        hit = next(r for r in rows if r["midi"] == 77)
        self.assertEqual(hit["provenance"], PROVENANCE_TRANSFERRED_EXTRAPOLATED)
        self.assertEqual(hit["extrapolation_distance_semitones"], 6)

    def test_legacy_scalar_is_combined_estimate(self):
        cfg = CalibrationConfig(transfer_method="legacy_target_calibrated")
        rows = build_final_rows(
            instrument="bass_clarinet",
            dynamic="mf",
            empirical={50: 10.0},
            transfer={72: 8.0},
            cfg=cfg,
            transfer_method="legacy_target_calibrated",
        )
        hit = next(r for r in rows if r["midi"] == 72)
        self.assertEqual(hit["provenance"], PROVENANCE_COMBINED_ESTIMATE)


class TestPrecedenceAndWeights(unittest.TestCase):
    def test_empirical_only_never_replaces_measurement(self):
        cfg = CalibrationConfig(combination_method="empirical_only")
        val, method, we, wt = combine_empirical_and_modelled(10.0, 99.0, cfg)
        self.assertEqual(val, 10.0)
        self.assertEqual(method, "empirical_only")
        self.assertEqual((we, wt), (1.0, 0.0))

    def test_equal_weight_legacy(self):
        cfg = CalibrationConfig(combination_method="legacy_equal_weight", empirical_weight=0.5)
        val, method, we, wt = combine_empirical_and_modelled(10.0, 20.0, cfg)
        self.assertAlmostEqual(val, 15.0)
        self.assertEqual((we, wt), (0.5, 0.5))

    def test_fixed_unequal_via_legacy_weights(self):
        cfg = CalibrationConfig(
            combination_method="legacy_equal_weight",
            empirical_weight=0.65,
            transfer_weight=0.35,
        )
        val, _, we, wt = combine_empirical_and_modelled(10.0, 20.0, cfg)
        self.assertAlmostEqual(val, 0.65 * 10 + 0.35 * 20)
        self.assertAlmostEqual(we + wt, 1.0)

    def test_one_zero_weights(self):
        cfg = CalibrationConfig(
            combination_method="legacy_equal_weight",
            empirical_weight=1.0,
            transfer_weight=0.0,
        )
        val, _, we, wt = combine_empirical_and_modelled(10.0, 20.0, cfg)
        self.assertAlmostEqual(val, 10.0)
        self.assertEqual((we, wt), (1.0, 0.0))
        cfg.empirical_weight = 0.0
        cfg.transfer_weight = 1.0
        val, _, we, wt = combine_empirical_and_modelled(10.0, 20.0, cfg)
        self.assertAlmostEqual(val, 20.0)

    def test_media_of_prefer_measured_keeps_iowa(self):
        iowa = build_layer("bass_clarinet", "IOWA", "ordinario", "mf", {60: 10.0}, "measured")
        orch = build_layer("bass_clarinet", "ORCH", "ordinario", "mf", {60: 99.0}, "orchidea_family_transfer_estimate")
        media = media_of([iowa, orch], prefer_measured=True)
        self.assertEqual(media.cells[60].value, 10.0)
        self.assertEqual(media.cells[60].origin, "measured")


class TestExtrapolationLimits(unittest.TestCase):
    def test_distances(self):
        self.assertEqual(extrapolation_distance_semitones(71, 37, 71), 0)
        self.assertEqual(extrapolation_distance_semitones(72, 37, 71), 1)
        self.assertEqual(extrapolation_distance_semitones(77, 37, 71), 6)
        self.assertEqual(extrapolation_distance_semitones(99, 37, 71), 28)

    def test_long_extrapolation_rejected(self):
        cfg = CalibrationConfig(allow_long_extrapolation=False)
        rows = build_final_rows(
            instrument="bass_clarinet",
            dynamic="pp",
            empirical={37: 8.0, 71: 9.0},
            transfer={99: 4.0},
            cfg=cfg,
        )
        hit = next(r for r in rows if r["midi"] == 99)
        self.assertEqual(hit["extrapolation_distance_semitones"], 28)
        self.assertEqual(hit["qa_flag"], "REJECTED")
        self.assertIsNone(hit["final_value"])


class TestTwoLogRatioNoLeakage(unittest.TestCase):
    def test_parameters_unchanged_when_target_overlap_removed_from_prediction_midis(self):
        iowa_cl = {m: 10.0 + 0.1 * (m - 50) for m in range(50, 80)}
        orch_cl = {m: 12.0 + 0.1 * (m - 50) for m in range(50, 80)}
        iowa_b = {m: 11.0 + 0.12 * (m - 50) for m in range(50, 80)}
        midis = list(range(50, 80))
        full = two_log_ratio_transfer(orch_cl, iowa_cl, iowa_b, midis)
        hidden = {m: v for m, v in iowa_b.items() if m < 50 or m > 77}
        # hide 50–77 from the *target* used to fit L_instr
        held = two_log_ratio_transfer(orch_cl, iowa_cl, hidden, midis)
        # L_coll is clarinet-only: unchanged if the target instrument is removed
        self.assertEqual(full.n_L_coll_anchors, held.n_L_coll_anchors)
        self.assertEqual(full.L_coll, held.L_coll)
        self.assertGreater(full.n_L_instr_anchors, held.n_L_instr_anchors)
        # removing the entire target collection leaves L_coll intact and L_instr UNDETERMINED
        none = two_log_ratio_transfer(orch_cl, iowa_cl, {}, midis)
        self.assertEqual(none.L_coll, full.L_coll)
        self.assertEqual(none.L_instr_status, "UNDETERMINED")
        self.assertEqual(none.curve, {})

    def test_holdout_ln_rmse_order(self):
        iowa_cl = {m: 20.0 * math.exp(-0.01 * (m - 55)) for m in range(40, 85)}
        orch_cl = {m: v * 1.1 for m, v in iowa_cl.items()}
        iowa_b = {m: v * 1.05 for m, v in iowa_cl.items() if 37 <= m <= 82}
        fit = {m: v for m, v in iowa_b.items() if m < 50 or m > 77}
        hold = {m: iowa_b[m] for m in iowa_b if 50 <= m <= 77}
        est = two_log_ratio_transfer(orch_cl, iowa_cl, fit, list(hold))
        mets = agreement_metrics(est.curve, hold)
        self.assertGreaterEqual(mets["n"], 5)
        self.assertLess(mets["ln_rmse"], 0.35)

    def test_iowa_excluded_from_validation_list(self):
        est = two_log_ratio_transfer({60: 10.0}, {60: 8.0}, {60: 9.0}, [60])
        self.assertIn("IOWA", est.validation_excluded)


class TestValidationOptimisation(unittest.TestCase):
    def test_obvious_weight(self):
        # truth = empirical exactly → optimal w = 1
        emp = {m: 10.0 + m for m in range(10)}
        trans = {m: 100.0 for m in range(10)}
        truth = dict(emp)
        got = optimal_weight(emp, trans, truth, min_n=5)
        self.assertIsNotNone(got)
        self.assertGreaterEqual(got["weight"], 0.95)

    def test_too_small_overlap_returns_none(self):
        self.assertIsNone(optimal_weight({1: 1.0}, {1: 2.0}, {1: 1.5}, min_n=5))


class TestMissingAndDeterminism(unittest.TestCase):
    def test_empty_and_nan(self):
        cfg = CalibrationConfig()
        self.assertEqual(combine_empirical_and_modelled(None, None, cfg)[0], None)
        self.assertEqual(combine_empirical_and_modelled(float("nan"), 3.0, cfg)[0], 3.0)
        self.assertEqual(agreement_metrics({}, {1: 1.0})["n"], 0)

    def test_determinism(self):
        iowa_cl = {m: 10.0 for m in range(50, 70)}
        orch_cl = {m: 11.0 for m in range(50, 70)}
        iowa_b = {m: 12.0 for m in range(50, 70)}
        a = two_log_ratio_transfer(orch_cl, iowa_cl, iowa_b, list(range(50, 70)))
        b = two_log_ratio_transfer(orch_cl, iowa_cl, iowa_b, list(range(50, 70)))
        self.assertEqual(a.curve, b.curve)

    def test_config_loads(self):
        cfg = load_config(ROOT / "calibration.yaml")
        self.assertEqual(cfg.combination_method, "empirical_only")
        self.assertEqual(cfg.transfer_method, "two_log_ratio")
        self.assertFalse(cfg.allow_review_required)
        self.assertEqual(cfg.transfer_field_mode, "pooled")
        self.assertEqual(cfg.transfer_field_weight, "anchors")
        self.assertEqual(cfg.transfer_field_min_collections, 2)


class TestGridAndAcousticLiterals(unittest.TestCase):
    def test_bass_clarinet_grid_and_ceiling(self):
        iowa = build_layer(
            "bass_clarinet",
            "IOWA",
            "ordinario",
            "mf",
            {37: 10.0, 82: 12.0},
            "measured",
        )
        phil = build_layer("bass_clarinet", "PHIL", "ordinario", "mf", {84: 9.0}, "measured")
        project_layers = {("IOWA", "mf"): iowa}
        # include phil via a fake IOWA-shaped dict for ceiling helper
        layers = {("IOWA", "mf"): iowa}
        grid = _note_grid(layers, "ordinario", "bass_clarinet")
        self.assertEqual(grid[0], 34)
        self.assertEqual(grid[-1], 82)
        self.assertEqual(measured_ceiling_midi(layers, "bass_clarinet"), 82)
        iowa.cells[84] = phil.cells[84]
        self.assertEqual(measured_ceiling_midi(layers, "bass_clarinet"), 84)

    def test_acoustic_table_has_literals(self):
        iowa = build_layer("bass_clarinet", "IOWA", "ordinario", "mf", {60: 10.0}, "measured")
        layers = {("IOWA", "mf"): iowa, ("ORCH", "mf"): build_layer("bass_clarinet", "ORCH", "ordinario", "mf", {60: 20.0}, "orchidea_family_transfer_estimate")}
        wb = Workbook()
        ws = wb.active
        _write_acoustic_table(ws, [60], layers, "ordinario", 82, "x.xlsx", "bass_clarinet", "Bass clarinet in Bb_Media")
        self.assertEqual(ws.cell(2, 5).value, 10.0)
        self.assertFalse(str(ws.cell(2, 5).value).startswith("=") if ws.cell(2, 5).value is not None else True)
        path = Path(tempfile.gettempdir()) / "ste_acoust_lit.xlsx"
        wb.save(path)
        raw = load_workbook(path, data_only=True)
        self.assertIsNotNone(raw.active.cell(2, 5).value)
        raw.close()


class TestAcceptancePolicy(unittest.TestCase):
    def _pp_span_rows(self, extra_transfer, cfg=None):
        cfg = cfg or CalibrationConfig()
        return build_final_rows(
            instrument="bass_clarinet",
            dynamic="pp",
            empirical={37: 8.0, 71: 9.0},
            transfer=extra_transfer,
            cfg=cfg,
        )

    def test_high_measured_is_accepted(self):
        rows = self._pp_span_rows({71: 99.0})
        hit = next(r for r in rows if r["midi"] == 71)
        self.assertEqual(hit["qa_flag"], "HIGH")
        self.assertTrue(hit["automatic_acceptance"])
        self.assertTrue(hit["accepted_final"])
        self.assertEqual(hit["validation_status"], "accepted")
        self.assertEqual(hit["final_value"], 9.0)

    def test_moderate_short_extrap_is_accepted(self):
        rows = self._pp_span_rows({72: 7.0, 73: 7.1, 74: 7.2})
        for midi, dist in ((72, 1), (73, 2), (74, 3)):
            hit = next(r for r in rows if r["midi"] == midi)
            self.assertEqual(hit["provenance"], PROVENANCE_TRANSFERRED_EXTRAPOLATED)
            self.assertEqual(hit["extrapolation_distance_semitones"], dist)
            self.assertEqual(hit["qa_flag"], "MODERATE")
            self.assertTrue(hit["accepted_final"])
            self.assertEqual(hit["validation_status"], "accepted")

    def test_review_required_not_accepted_by_default(self):
        rows = self._pp_span_rows({75: 6.0, 76: 6.1, 77: 6.2})
        for midi, dist in ((75, 4), (76, 5), (77, 6)):
            hit = next(r for r in rows if r["midi"] == midi)
            self.assertEqual(hit["qa_flag"], "REVIEW_REQUIRED")
            self.assertFalse(hit["automatic_acceptance"])
            self.assertFalse(hit["acceptance_override"])
            self.assertFalse(hit["accepted_final"])
            self.assertEqual(hit["validation_status"], "review_required")
            self.assertIsNotNone(hit["final_value"])

    def test_rejected_not_accepted(self):
        rows = self._pp_span_rows({99: 4.0})
        hit = next(r for r in rows if r["midi"] == 99)
        self.assertEqual(hit["qa_flag"], "REJECTED")
        self.assertFalse(hit["accepted_final"])
        self.assertEqual(hit["validation_status"], "rejected")
        self.assertIsNone(hit["final_value"])

    def test_review_override_preserves_qa_flag(self):
        cfg = CalibrationConfig(allow_review_required=True)
        rows = self._pp_span_rows({75: 6.0}, cfg)
        hit = next(r for r in rows if r["midi"] == 75)
        self.assertEqual(hit["qa_flag"], "REVIEW_REQUIRED")
        self.assertFalse(hit["automatic_acceptance"])
        self.assertTrue(hit["acceptance_override"])
        self.assertTrue(hit["accepted_final"])
        self.assertEqual(hit["validation_status"], "accepted")

    def test_acoustic_table_follows_final_calibration(self):
        cfg = CalibrationConfig()
        rows = self._pp_span_rows({m: 7.0 for m in range(72, 78)})
        idx = qa_index(rows)
        iowa = build_layer("bass_clarinet", "IOWA", "ordinario", "pp", {37: 8.0, 71: 9.0}, "measured")
        orch = build_layer(
            "bass_clarinet",
            "ORCH",
            "ordinario",
            "pp",
            {m: 7.0 for m in range(72, 78)},
            "orchidea_family_transfer_estimate",
        )
        layers = {("IOWA", "pp"): iowa, ("ORCH", "pp"): orch}
        wb = Workbook()
        _write_acoustic_table(
            wb.active,
            list(range(37, 78)),
            layers,
            "ordinario",
            82,
            "x.xlsx",
            "bass_clarinet",
            "Bass clarinet in Bb_Media",
            final_by_cell=idx,
        )
        statuses = {}
        for rec in wb.active.iter_rows(min_row=2, values_only=True):
            midi, dyn, status = rec[2], rec[3], rec[14]
            if dyn == "pp":
                statuses[int(midi)] = status
        self.assertEqual(statuses[72], "accepted")
        self.assertEqual(statuses[74], "accepted")
        self.assertEqual(statuses[75], "review_required")
        self.assertEqual(statuses[77], "review_required")

    def test_acceptance_counts_reconcile(self):
        rows = self._pp_span_rows({72: 7.0, 75: 6.0, 99: 4.0})
        counts = acceptance_counts(rows)
        self.assertEqual(counts["rows_total"], len(rows))
        self.assertEqual(
            counts["rows_accepted"] + counts["rows_review_required"] + counts["rows_rejected"],
            counts["rows_total"],
        )
        self.assertGreaterEqual(counts["rows_review_required"], 1)
        self.assertGreaterEqual(counts["rows_rejected"], 1)


class TestLnRmseAndShape(unittest.TestCase):
    def test_ln_rmse_applies_square_root(self):
        pred = {1: math.e ** 2, 2: math.e ** 3}
        truth = {1: math.e, 2: math.e}
        mets = agreement_metrics(pred, truth)
        mse = ((1.0) ** 2 + (2.0) ** 2) / 2.0
        self.assertAlmostEqual(mets["ln_rmse"], math.sqrt(mse))
        self.assertNotAlmostEqual(mets["ln_rmse"], mse)

    def test_shape_metrics_analytical_fixture(self):
        truth = {m: 10.0 * math.exp(-0.02 * (m - 50)) for m in range(50, 60)}
        pred = {m: v / 2.0 for m, v in truth.items()}
        shape = shape_normalised_metrics(pred, truth)
        self.assertAlmostEqual(shape["collection_log_offset"], math.log(2.0))
        self.assertAlmostEqual(shape["collection_scale_factor"], 2.0)
        self.assertAlmostEqual(shape["shape_ln_rmse"], 0.0)
        self.assertAlmostEqual(shape["shape_ln_mae"], 0.0)
        self.assertEqual(shape["collection_scale_factor"], math.exp(shape["collection_log_offset"]))

    def test_shape_does_not_change_production_values(self):
        iowa = {60: 10.0, 61: 11.0}
        transfer = {60: 20.0, 61: 22.0}
        before = dict(iowa)
        shape_normalised_metrics(transfer, iowa)
        self.assertEqual(iowa, before)

    def test_iowa_excluded_from_transfer_validation_tables(self):
        from ste_lab.calibration import TransferEstimate

        measured = {
            ("IOWA", "mf"): {60: 10.0, 61: 11.0},
            ("PHIL", "mf"): {60: 12.0, 61: 13.0},
        }
        transfer = {"mf": TransferEstimate(curve={60: 9.0, 61: 9.5}, method="two_log_ratio", validation_excluded=["IOWA"])}
        rows, _offsets, excluded = build_validation_tables(measured, transfer, CalibrationConfig())
        self.assertIn("IOWA", excluded)
        self.assertTrue(all(r["validation_collection"] != "IOWA" for r in rows))
        self.assertTrue(any(r.get("shape_ln_rmse") is not None for r in rows))


class TestMeasuredDataProvenance(unittest.TestCase):
    def test_family_donor_map(self):
        self.assertEqual(family_donor_instrument("bass_clarinet"), "clarinet")
        self.assertEqual(family_donor_instrument("piccolo"), "flute")
        self.assertEqual(family_donor_instrument("english_horn"), "oboe")
        self.assertEqual(family_donor_instrument("contrabassoon"), "bassoon")
        self.assertIsNone(family_donor_instrument("clarinet"))
        self.assertIsNone(family_donor_instrument("viola"))

    def test_woodwind_orch_is_transferred_not_measured(self):
        wb = Workbook()
        ws = wb.active
        write_measured_data(
            ws,
            {
                ("IOWA", "mf"): {60: 10.0},
                ("ORCH", "mf"): {60: 12.0},
                ("PHIL", "mf"): {60: 11.0},
            },
            "bass_clarinet",
            origins={
                "IOWA": "measured",
                "ORCH": "orchidea_family_transfer_estimate",
                "PHIL": "measured",
            },
        )
        headers = [c.value for c in ws[1]]
        rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
        by_coll = {r["collection"]: r for r in rows}
        self.assertEqual(by_coll["IOWA"]["provenance"], PROVENANCE_MEASURED)
        self.assertEqual(by_coll["IOWA"]["source_type"], "observation")
        self.assertEqual(by_coll["IOWA"]["source_instrument"], "bass_clarinet")
        self.assertEqual(by_coll["ORCH"]["provenance"], PROVENANCE_TRANSFERRED)
        self.assertEqual(by_coll["ORCH"]["source_instrument"], "clarinet")
        self.assertNotEqual(by_coll["ORCH"]["source_type"], "observation")
        self.assertEqual(by_coll["PHIL"]["provenance"], PROVENANCE_MEASURED)

    def test_woodwind_orch_fallback_without_origins(self):
        wb = Workbook()
        write_measured_data(wb.active, {("ORCH", "mf"): {60: 12.0}}, "bass_clarinet")
        headers = [c.value for c in wb.active[1]]
        rec = dict(zip(headers, next(wb.active.iter_rows(min_row=2, values_only=True))))
        self.assertEqual(rec["provenance"], PROVENANCE_TRANSFERRED)
        self.assertEqual(rec["source_instrument"], "clarinet")


if __name__ == "__main__":
    unittest.main()
