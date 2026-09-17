# -*- coding: utf-8 -*-
"""Synthetic fixtures for 1.5.5 donor-dynamic, provenance and invariance fixes.

Isolated TemporaryDirectory outputs only. Does not touch historical recordings,
input workbooks, preparatory workbooks, or research analyses.
"""
from __future__ import annotations

import io
import math
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import pandas as pd
from openpyxl import Workbook

from build_ste_preparatory import build
from run_ste_from_preparatory import load_preparatory, run
from ste_lab.calibration import CalibrationConfig
from ste_lab.notes import midi_to_label
from ste_lab.relations import (
    L_INVARIANCE_HOLD,
    Relation,
    closest_production_relation,
    collection_from_grade,
    l_invariance_status,
    production_relations_from_effects,
)
from ste_lab.transfer import Anchor


NOTES = {60: "C4", 62: "D4", 64: "E4"}


def _exp_clip(L: float) -> float:
    return math.exp(min(3.0, max(-3.0, L)))


def _value_at(df: pd.DataFrame, midi: int) -> float:
    hit = df.loc[df["MIDI"] == midi]
    if hit.empty:
        raise AssertionError(f"MIDI {midi} missing")
    return float(hit.iloc[0]["Combined density metric"])


def _layer_table(path: Path, sheet: str) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=sheet)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def _write_arco(wb: Workbook, iowa_pp: float = 20.0) -> None:
    arco = wb.active
    arco.title = "Paste_arco"
    arco.append(["note", "midi", "IOWA_pp", "IOWA_mf", "IOWA_ff", "ORCH_pp", "ORCH_mf", "ORCH_ff"])
    for midi, lab in NOTES.items():
        shift = midi - 60
        arco.append(
            [
                lab,
                midi,
                iowa_pp + shift,
                22.0 + shift,
                24.0 + shift,
                10.0 + 0.5 * shift,
                11.0 + 0.5 * shift,
                12.0 + 0.5 * shift,
            ]
        )


def _write_effects(wb: Workbook, recorded: dict[int, float] | None = None) -> None:
    eff = wb.create_sheet("Paste_effects_mf")
    headers = ["note", "midi"]
    if recorded is not None:
        headers.append("ORCH_ponticello_mf")
    eff.append(headers)
    for midi, lab in NOTES.items():
        row = [lab, midi]
        if recorded is not None:
            row.append(recorded.get(midi))
        eff.append(row)


def _anchor_row(effect, grade, dyn, midi, src, tgt, collection="", source="ordinario"):
    return [
        effect,
        grade,
        collection,
        dyn,
        source,
        effect,
        NOTES[midi],
        midi,
        src,
        tgt,
    ]


def _write_prep(
    path: Path,
    rows: list[list],
    *,
    include_collection: bool = True,
    recorded: dict[int, float] | None = None,
    iowa_pp: float = 20.0,
) -> None:
    wb = Workbook()
    _write_arco(wb, iowa_pp=iowa_pp)
    _write_effects(wb, recorded=recorded)
    anc = wb.create_sheet("Anchors_all")
    headers = [
        "effect",
        "evidence_grade",
        "collection",
        "dynamic",
        "source_cond",
        "target_cond",
        "note",
        "midi",
        "sourceCDM",
        "targetCDM",
    ]
    if not include_collection:
        headers = [h for h in headers if h != "collection"]
        rows = [r[:2] + r[3:] for r in rows]
    anc.append(headers)
    for row in rows:
        anc.append(row)
    wb.save(path)


def _ratio_rows(effect: str, grade: str, dyn: str, ratio: float, collection: str, src0: float = 10.0):
    rows = []
    for midi in NOTES:
        src = src0 + 0.5 * (midi - 60)
        rows.append(_anchor_row(effect, grade, dyn, midi, src, src * ratio, collection=collection))
    return rows


def _write_compiled(path: Path, curve: dict[int, float]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Research_Core"
    ws.append(["note", "spectral_mass"])
    for midi, value in sorted(curve.items()):
        ws.append([midi_to_label(midi), float(value)])
    wb.save(path)


def _stable(tree: Path, *parts: str, curve: dict[int, float]) -> None:
    folder = tree.joinpath(*parts, "_Sustains_Stable", "spectral_analyser")
    name = f"{'_'.join(parts)}_compiled_density_metrics_research.xlsx"
    _write_compiled(folder / name, curve)


class TestCollectionFromGrade(unittest.TestCase):
    def test_unambiguous_tokens(self):
        self.assertEqual(collection_from_grade("Empirical_ORCH"), "ORCH")
        self.assertEqual(collection_from_grade("prediction_Philharmonia"), "PHIL")
        self.assertEqual(collection_from_grade("prediction_McGill"), "MCGILL")

    def test_ambiguous_grade_is_not_mcgill(self):
        self.assertIsNone(collection_from_grade("prediction"))
        self.assertIsNone(collection_from_grade("invented"))


class TestLoadKeepsEligibleDonors(unittest.TestCase):
    def test_p_is_kept_when_mf_exists(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prep.xlsx"
            rows = _ratio_rows("sul ponticello", "prediction_Philharmonia", "p", 0.5, "PHIL")
            rows += _ratio_rows("sul ponticello", "prediction_Philharmonia", "mf", 0.8, "PHIL")
            _write_prep(path, rows)
            pack = load_preparatory(path)
            spec = next(e for e in pack["effects"] if e["name"] == "sul ponticello")
            self.assertIn("p", spec["anchors_by_dyn"])
            self.assertIn("mf", spec["anchors_by_dyn"])
            self.assertEqual(spec["collection"], "PHIL")

    def test_mixed_sets_keep_both_dynamics(self):
        cases = (("p", "mf"), ("mf", "f"), ("p", "f"))
        for left, right in cases:
            with self.subTest(pair=f"{left}+{right}"):
                with tempfile.TemporaryDirectory() as tmp:
                    path = Path(tmp) / "prep.xlsx"
                    rows = _ratio_rows("sul ponticello", "prediction_Philharmonia", left, 0.5, "PHIL")
                    rows += _ratio_rows("sul ponticello", "prediction_Philharmonia", right, 0.8, "PHIL")
                    _write_prep(path, rows)
                    pack = load_preparatory(path)
                    spec = next(e for e in pack["effects"] if e["name"] == "sul ponticello")
                    self.assertEqual(set(spec["anchors_by_dyn"]), {left, right})


class TestClosestDonorExport(unittest.TestCase):
    def test_p_plus_mf_exports_10_not_16(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prep.xlsx"
            out = Path(tmp) / "out"
            rows = _ratio_rows("sul ponticello", "prediction_Philharmonia", "p", 0.5, "PHIL")
            rows += _ratio_rows("sul ponticello", "prediction_Philharmonia", "mf", 0.8, "PHIL")
            _write_prep(path, rows)
            single = CalibrationConfig(transfer_field_mode="single")
            with patch("run_ste_from_preparatory.load_config", return_value=single):
                written = run(path, out, instrument="violin")
            ste = out / "Violin_STE_sul_ponticello_IOWA_ORCH.xlsx"
            self.assertTrue(ste.exists(), written)
            iowa_pp = _layer_table(ste, "violin_IOWA_pp")
            self.assertAlmostEqual(_value_at(iowa_pp, 60), 10.0)
            self.assertNotAlmostEqual(_value_at(iowa_pp, 60), 16.0)
            em = pd.read_excel(ste, sheet_name="Evidence_Map")
            hit = em[(em["collection"].astype(str).str.upper() == "IOWA") & (em["dynamic"].astype(str) == "pp")]
            self.assertEqual(str(hit.iloc[0]["l_donor_dynamic"]).strip().lower(), "p")
            zen = out / "Violin_Zenodo_collections_sul_ponticello.xlsx"
            self.assertTrue(zen.exists())
            media = pd.read_excel(zen, sheet_name="Violin_Media")
            row = media.loc[media["Note"].astype(str) == "C4"]
            self.assertFalse(row.empty)
            iowa = float(row.iloc[0]["IOWA pp"])
            orch = float(row.iloc[0]["ORCH pp"])
            self.assertAlmostEqual(iowa, 10.0)
            self.assertAlmostEqual(float(row.iloc[0]["Media pp"]), 0.5 * (iowa + orch))
            self.assertTrue(math.isfinite(iowa) and math.isfinite(orch))

    def test_mf_plus_f_and_p_plus_f_donor_labels(self):
        expect = {
            ("mf", "f"): {"pp": "mf", "mf": "mf", "ff": "f"},
            ("p", "f"): {"pp": "p", "mf": "f", "ff": "f"},
        }
        ratios = {"p": 0.5, "mf": 0.8, "f": 0.4}
        for pair, donors in expect.items():
            with self.subTest(pair="+".join(pair)):
                with tempfile.TemporaryDirectory() as tmp:
                    path = Path(tmp) / "prep.xlsx"
                    out = Path(tmp) / "out"
                    rows = []
                    for dyn in pair:
                        rows += _ratio_rows("sul ponticello", "prediction_Philharmonia", dyn, ratios[dyn], "PHIL")
                    _write_prep(path, rows)
                    single = CalibrationConfig(transfer_field_mode="single")
                    with patch("run_ste_from_preparatory.load_config", return_value=single):
                        run(path, out, instrument="violin")
                    ste = out / "Violin_STE_sul_ponticello_IOWA_ORCH.xlsx"
                    em = pd.read_excel(ste, sheet_name="Evidence_Map")
                    iowa_pp = _layer_table(ste, "violin_IOWA_pp")
                    donor_pp = donors["pp"]
                    self.assertAlmostEqual(_value_at(iowa_pp, 60), 20.0 * ratios[donor_pp])
                    for dyn, donor in donors.items():
                        hit = em[
                            (em["collection"].astype(str).str.upper() == "IOWA")
                            & (em["dynamic"].astype(str) == dyn)
                        ]
                        self.assertFalse(hit.empty)
                        raw = hit.iloc[0]["l_donor_dynamic"]
                        got = "" if pd.isna(raw) else str(raw).strip().lower()
                        if donor == dyn:
                            self.assertIn(got, {"", dyn})
                        else:
                            self.assertEqual(got, donor)


class TestRelationProvenance(unittest.TestCase):
    def test_phil_ff_and_orch_mf_keep_own_collection(self):
        for reverse in (False, True):
            with self.subTest(reverse=reverse):
                with tempfile.TemporaryDirectory() as tmp:
                    path = Path(tmp) / "prep.xlsx"
                    out = Path(tmp) / "out"
                    orch = _ratio_rows("sul ponticello", "Empirical_ORCH", "mf", 0.8, "ORCH")
                    phil = _ratio_rows("sul ponticello", "prediction_Philharmonia", "ff", 0.5, "PHIL")
                    rows = phil + orch if reverse else orch + phil
                    recorded = {midi: 11.0 * 0.8 + 0.4 * (midi - 60) for midi in NOTES}
                    _write_prep(path, rows, recorded=recorded)
                    pack = load_preparatory(path)
                    specs = [e for e in pack["effects"] if e["name"] == "sul ponticello"]
                    colls = {e["collection"] for e in specs}
                    self.assertEqual(colls, {"ORCH", "PHIL"})
                    rels = production_relations_from_effects(pack["effects"], "violin")
                    by_dyn = {r.dynamic: r.collection for r in rels}
                    self.assertEqual(by_dyn["mf"], "ORCH")
                    self.assertEqual(by_dyn["ff"], "PHIL")
                    single = CalibrationConfig(transfer_field_mode="single")
                    with patch("run_ste_from_preparatory.load_config", return_value=single):
                        run(path, out, instrument="violin")
                    ste = out / "Violin_STE_sul_ponticello_IOWA_ORCH.xlsx"
                    orch_mf = _layer_table(ste, "violin_ORCH_mf")
                    self.assertAlmostEqual(_value_at(orch_mf, 60), recorded[60])
                    iowa_pp = _layer_table(ste, "violin_IOWA_pp")
                    self.assertAlmostEqual(_value_at(iowa_pp, 60), 16.0)
                    em = pd.read_excel(ste, sheet_name="Evidence_Map")
                    pp = em[(em["collection"].astype(str).str.upper() == "IOWA") & (em["dynamic"].astype(str) == "pp")]
                    self.assertEqual(str(pp.iloc[0]["l_donor_dynamic"]).strip().lower(), "mf")
                    zen = out / "Violin_Zenodo_collections_sul_ponticello.xlsx"
                    media = pd.read_excel(zen, sheet_name="Violin_Media")
                    row = media.loc[media["Note"].astype(str) == "C4"].iloc[0]
                    self.assertAlmostEqual(float(row["ORCH mf"]), recorded[60])
                    self.assertAlmostEqual(
                        float(row["Media mf"]),
                        0.5 * (float(row["IOWA mf"]) + float(row["ORCH mf"])),
                    )

    def test_same_dynamic_two_collections_pool_L_not_absolute_cdm(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prep.xlsx"
            out = Path(tmp) / "out"
            rows = _ratio_rows("sul ponticello", "Empirical_ORCH", "mf", 0.8, "ORCH")
            rows += _ratio_rows("sul ponticello", "prediction_Philharmonia", "mf", 0.5, "PHIL")
            _write_prep(path, rows)
            pooled = CalibrationConfig(transfer_field_mode="pooled", transfer_field_min_collections=2)
            with patch("run_ste_from_preparatory.load_config", return_value=pooled):
                run(path, out, instrument="violin")
            ste = out / "Violin_STE_sul_ponticello_IOWA_ORCH.xlsx"
            iowa_mf = _layer_table(ste, "violin_IOWA_mf")
            L_pool = 0.5 * (math.log(0.8) + math.log(0.5))
            self.assertAlmostEqual(_value_at(iowa_mf, 60), 22.0 * _exp_clip(L_pool))
            single = CalibrationConfig(transfer_field_mode="single")
            with patch("run_ste_from_preparatory.load_config", return_value=single):
                run(path, out / "single", instrument="violin")
            iowa_single = _layer_table(out / "single" / "Violin_STE_sul_ponticello_IOWA_ORCH.xlsx", "violin_IOWA_mf")
            chosen = closest_production_relation(
                production_relations_from_effects(load_preparatory(path)["effects"], "violin"),
                "mf",
            )
            self.assertEqual(chosen.collection, "ORCH")
            self.assertAlmostEqual(_value_at(iowa_single, 60), 22.0 * 0.8)

    def test_legacy_grade_inference_and_ambiguous_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "legacy.xlsx"
            orch = _ratio_rows("sul ponticello", "Empirical_ORCH", "mf", 0.8, "")
            phil = _ratio_rows("sul ponticello", "prediction_Philharmonia", "ff", 0.5, "")
            _write_prep(path, orch + phil, include_collection=False)
            pack = load_preparatory(path)
            specs = [e for e in pack["effects"] if e["name"] == "sul ponticello"]
            self.assertEqual({e["collection"] for e in specs}, {"ORCH", "PHIL"})
            amb = Path(tmp) / "amb.xlsx"
            rows = _ratio_rows("sul ponticello", "prediction", "mf", 0.8, "")
            buf = io.StringIO()
            with redirect_stdout(buf):
                _write_prep(amb, rows, include_collection=False)
                pack_amb = load_preparatory(amb)
            spec = pack_amb["effects"][0]
            self.assertEqual(spec["collection"], "")
            self.assertTrue(spec["ambiguous_collection"])
            self.assertIn("AMBIGUOUS collection", buf.getvalue())


class TestInvarianceDiagnostic(unittest.TestCase):
    def _rel(self, dyn: str, targets: list[float], collection: str = "PHIL", midis=None) -> Relation:
        midis = midis or [60, 62, 64]
        anchors = [Anchor(m, 1.0, math.exp(t)) for m, t in zip(midis, targets)]
        return Relation(
            kind="technique",
            instrument="violin",
            source_cond="ordinario",
            target_cond="sul ponticello",
            collection=collection,
            dynamic=dyn,
            anchors=anchors,
        )

    def test_crossing_curves_identical_means_are_not_held(self):
        a = self._rel("p", [-1.0, 0.0, 1.0])
        b = self._rel("mf", [1.0, 0.0, -1.0])
        report = l_invariance_status([a, b])
        self.assertEqual(report.status, "compared")
        self.assertAlmostEqual(report.mean_L_spread, 0.0)
        self.assertAlmostEqual(report.max_abs_delta, 2.0)
        self.assertEqual(report.heuristic, "wide")
        self.assertEqual(report.n_shared, 3)
        self.assertEqual(report.collection, "PHIL")
        self.assertIn("p", report.pair)
        self.assertIn("mf", report.pair)
        self.assertNotEqual(report.status, "held")

    def test_identical_curves(self):
        shared = [-0.2, -0.2, -0.2]
        report = l_invariance_status([self._rel("p", shared), self._rel("mf", shared)])
        self.assertEqual(report.status, "compared")
        self.assertEqual(report.heuristic, "held")
        self.assertAlmostEqual(report.max_abs_delta, 0.0)
        self.assertLessEqual(report.max_abs_delta, L_INVARIANCE_HOLD)

    def test_partial_overlap(self):
        a = self._rel("p", [-0.1, -0.1], midis=[60, 62])
        b = self._rel("mf", [0.5, 0.5], midis=[62, 64])
        report = l_invariance_status([a, b])
        self.assertEqual(report.status, "compared")
        self.assertEqual(report.n_shared, 1)
        self.assertAlmostEqual(report.max_abs_delta, 0.6)

    def test_no_overlap(self):
        a = self._rel("p", [-0.1, -0.1], midis=[60, 62])
        b = self._rel("mf", [0.5, 0.5], midis=[72, 74])
        report = l_invariance_status([a, b])
        self.assertEqual(report.status, "insufficient_overlap")
        self.assertEqual(report.n_shared, 0)
        self.assertEqual(report.heuristic, "n/a")

    def test_single_donor_dynamic(self):
        report = l_invariance_status([self._rel("p", [-0.2, -0.2, -0.2])])
        self.assertEqual(report.status, "single_dynamic")

    def test_mixed_collections_are_not_compared_as_one_curve(self):
        a = self._rel("p", [-1.0, 0.0, 1.0], collection="PHIL")
        b = self._rel("mf", [1.0, 0.0, -1.0], collection="ORCH")
        report = l_invariance_status([a, b])
        self.assertEqual(report.status, "mixed_collection")
        self.assertEqual(report.heuristic, "n/a")


class TestBuilderRunnerIntegration(unittest.TestCase):
    def test_builder_keeps_p_and_exports_nearest_donor(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "deposit"
            generated = Path(tmp) / "generated"
            exports = Path(tmp) / "exports"
            generated.mkdir()
            midis = NOTES
            iowa = {midi: 20.0 + (midi - 60) for midi in midis}
            orch = {midi: 10.0 + 0.5 * (midi - 60) for midi in midis}
            wb = Workbook()
            media = wb.active
            media.title = "Violin_Media"
            media.append(["Note", "IOWA pp", "IOWA mf", "IOWA ff", "ORCH pp", "ORCH mf", "ORCH ff"])
            for midi, lab in midis.items():
                media.append([lab, iowa[midi], iowa[midi] + 2, iowa[midi] + 4, orch[midi], orch[midi] + 1, orch[midi] + 2])
            root.mkdir()
            wb.save(root / "Violin_Zenodo_collections_Arco_normal.xlsx")

            orch_tree = root / "Orchidea_violin"
            _stable(orch_tree, "ordinario", "pianissimo", curve=orch)
            _stable(orch_tree, "ordinario", "piano", curve=orch)
            _stable(orch_tree, "ordinario", "mezzo-forte", curve={m: v + 1 for m, v in orch.items()})
            _stable(orch_tree, "ordinario", "fortissimo", curve={m: v + 2 for m, v in orch.items()})
            _stable(orch_tree, "sul-ponticello", "piano", curve={m: 0.5 * v for m, v in orch.items()})
            _stable(
                orch_tree,
                "sul-ponticello",
                "mezzo-forte",
                curve={m: 0.8 * (v + 1) for m, v in orch.items()},
            )

            phil = root / "Philharmonia_violin"
            _stable(phil, "arco-normal", "piano", curve=orch)
            _stable(phil, "arco-normal", "fortissimo", curve={m: v + 2 for m, v in orch.items()})
            _stable(phil, "arco-sul-tasto", "piano", curve={m: 0.6 * v for m, v in orch.items()})

            prep = generated / "Violin_STE_preparatory.xlsx"
            buf = io.StringIO()
            with redirect_stdout(buf):
                build(root, instrument="violin", out_path=prep)
            self.assertTrue(prep.exists())
            pack = load_preparatory(prep)
            pont = next(e for e in pack["effects"] if e["name"] == "sul ponticello")
            self.assertIn("p", pont["anchors_by_dyn"])
            self.assertIn("mf", pont["anchors_by_dyn"])
            self.assertEqual(pont["collection"], "ORCH")
            tasto = next(e for e in pack["effects"] if e["name"] == "sul tasto")
            self.assertEqual(tasto["collection"], "PHIL")
            self.assertNotIn("ordinario", [e["name"] for e in pack["effects"]])

            single = CalibrationConfig(transfer_field_mode="single")
            with patch("run_ste_from_preparatory.load_config", return_value=single):
                with redirect_stdout(buf):
                    written = run(prep, exports, instrument="violin")
            ste = exports / "Violin_STE_sul_ponticello_IOWA_ORCH.xlsx"
            self.assertTrue(ste.exists(), written)
            iowa_pp = _layer_table(ste, "violin_IOWA_pp")
            self.assertAlmostEqual(_value_at(iowa_pp, 60), 10.0)
            em = pd.read_excel(ste, sheet_name="Evidence_Map")
            hit = em[(em["collection"].astype(str).str.upper() == "IOWA") & (em["dynamic"].astype(str) == "pp")]
            self.assertEqual(str(hit.iloc[0]["l_donor_dynamic"]).strip().lower(), "p")
            self.assertEqual(str(hit.iloc[0]["l_invariance"]).strip().lower(), "compared")
            zen = exports / "Violin_Zenodo_collections_sul_ponticello.xlsx"
            self.assertTrue(zen.exists())
            media_out = pd.read_excel(zen, sheet_name="Violin_Media")
            row = media_out.loc[media_out["Note"].astype(str) == "C4"].iloc[0]
            self.assertAlmostEqual(float(row["IOWA pp"]), 10.0)
            self.assertAlmostEqual(float(row["Media pp"]), 0.5 * (float(row["IOWA pp"]) + float(row["ORCH pp"])))
            for col in ("IOWA pp", "ORCH pp", "Media pp"):
                self.assertTrue(math.isfinite(float(row[col])))
            tasto_ste = exports / "Violin_STE_sul_tasto_IOWA_ORCH.xlsx"
            self.assertTrue(tasto_ste.exists())


class TestUnaffectedEmpiricalPath(unittest.TestCase):
    def test_mf_only_empirical_still_uses_mf_ratio(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prep.xlsx"
            out = Path(tmp) / "out"
            rows = _ratio_rows("sul ponticello", "Empirical_ORCH", "mf", 0.8, "ORCH")
            recorded = {midi: 11.0 * 0.8 + 0.4 * (midi - 60) for midi in NOTES}
            _write_prep(path, rows, recorded=recorded)
            single = CalibrationConfig(transfer_field_mode="single")
            with patch("run_ste_from_preparatory.load_config", return_value=single):
                run(path, out, instrument="violin")
            ste = out / "Violin_STE_sul_ponticello_IOWA_ORCH.xlsx"
            self.assertAlmostEqual(_value_at(_layer_table(ste, "violin_IOWA_pp"), 60), 16.0)
            self.assertAlmostEqual(_value_at(_layer_table(ste, "violin_ORCH_mf"), 60), recorded[60])


if __name__ == "__main__":
    unittest.main()
