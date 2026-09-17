# -*- coding: utf-8 -*-
"""Regression tests for Media export, combination provenance, config, and Fill."""
from __future__ import annotations

import inspect
import math
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from run_ste_effects_batch import build_layer
from ste_lab.calibration import (
    PROVENANCE_COMBINED_ESTIMATE,
    PROVENANCE_MEASURED,
    QA_HIGH,
    QA_MODERATE,
    QA_REVIEW,
    CalibrationConfig,
    ConfigurationError,
    build_final_rows,
    combine_empirical_and_modelled,
    load_config,
    remember_transfer_estimates,
    validate_combination_config,
)
from ste_lab.catalog import GENERATED_ORIGINS, normalize_origin
from ste_lab.gui_app import STEApp
from ste_lab.media_policy import (
    production_average_for_layer,
    production_media_of,
    resolve_iowa_orch_pair,
)
from ste_lab.session import Project
from ste_lab.transfer import InsufficientFillData, fill_missing, media_of
from ste_lab.zenodo_export import (
    _media_sheet_formula,
    export_zenodo_workbook,
    media_sheet_name,
)


def _independent_mean(values: list[float]) -> float:
    return sum(values) / len(values)


def _independent_mix(empirical: float, transfer: float, w_e: float) -> float:
    return w_e * empirical + (1.0 - w_e) * transfer


class TestStringMediaPolicy(unittest.TestCase):
    def test_both_present_are_arithmetic_mean_everywhere(self):
        expected = _independent_mean([10.0, 30.0])
        self.assertEqual(expected, 20.0)
        value, method, w_e, w_t = resolve_iowa_orch_pair(
            10.0, 30.0, instrument="violin", technique="con sordino"
        )
        self.assertEqual(value, expected)
        self.assertEqual(method, "string_mean")
        self.assertEqual((w_e, w_t), (0.5, 0.5))

        iowa = build_layer("violin", "IOWA", "con sordino", "mf", {60: 10.0}, "measured")
        orch = build_layer("violin", "ORCH", "con sordino", "mf", {60: 30.0}, "measured")
        phil = build_layer("violin", "Philharmonia", "con sordino", "mf", {60: 99.0}, "measured")
        media = production_media_of(
            [iowa, orch, phil],
            instrument="violin",
            technique="con sordino",
        )
        self.assertEqual(media.cells[60].value, expected)
        self.assertNotIn("99", str(media.cells[60].value))

    def test_one_present_is_not_halved(self):
        only_iowa, _, w_e, w_t = resolve_iowa_orch_pair(
            10.0, None, instrument="violin", technique="ordinario"
        )
        only_orch, _, w_e2, w_t2 = resolve_iowa_orch_pair(
            None, 30.0, instrument="violin", technique="ordinario"
        )
        self.assertEqual(only_iowa, 10.0)
        self.assertEqual((w_e, w_t), (1.0, 0.0))
        self.assertEqual(only_orch, 30.0)
        self.assertEqual((w_e2, w_t2), (0.0, 1.0))

    def test_neither_present_is_missing(self):
        value, _, _, _ = resolve_iowa_orch_pair(
            None, None, instrument="violin", technique="ordinario"
        )
        self.assertIsNone(value)


class TestWoodwindCombination(unittest.TestCase):
    def test_empirical_only_keeps_target(self):
        cfg = CalibrationConfig(combination_method="empirical_only")
        expected = 10.0
        val, method, we, wt = combine_empirical_and_modelled(10.0, 30.0, cfg)
        self.assertEqual(val, expected)
        self.assertEqual(method, "empirical_only")
        self.assertEqual((we, wt), (1.0, 0.0))

    def test_equal_weight_is_independent_midpoint(self):
        cfg = CalibrationConfig(combination_method="equal_weight")
        expected = _independent_mix(10.0, 30.0, 0.5)
        self.assertEqual(expected, 20.0)
        val, method, we, wt = combine_empirical_and_modelled(10.0, 30.0, cfg)
        self.assertEqual(val, expected)
        self.assertEqual(method, "equal_weight")
        self.assertEqual((we, wt), (0.5, 0.5))

    def test_legacy_equal_weight_uses_one_minus_empirical(self):
        cfg = CalibrationConfig(
            combination_method="legacy_equal_weight",
            empirical_weight=0.25,
            transfer_weight=0.75,
        )
        expected = _independent_mix(10.0, 30.0, 0.25)
        self.assertEqual(expected, 25.0)
        val, method, we, wt = combine_empirical_and_modelled(10.0, 30.0, cfg)
        self.assertEqual(val, expected)
        self.assertEqual(method, "legacy_equal_weight")
        self.assertEqual((we, wt), (0.25, 0.75))

    def test_missing_empirical_uses_transfer(self):
        cfg = CalibrationConfig(combination_method="empirical_only")
        val, method, we, wt = combine_empirical_and_modelled(None, 30.0, cfg)
        self.assertEqual(val, 30.0)
        self.assertEqual(method, "transfer_only")
        self.assertEqual((we, wt), (0.0, 1.0))

    def test_missing_both_stays_missing(self):
        cfg = CalibrationConfig(combination_method="equal_weight")
        val, _, _, _ = combine_empirical_and_modelled(None, None, cfg)
        self.assertIsNone(val)

    def test_rejected_transfer_is_not_promoted(self):
        cfg = CalibrationConfig()
        rows = build_final_rows(
            instrument="bass_clarinet",
            dynamic="mf",
            empirical={50: 10.0},
            transfer={90: 4.0},
            cfg=cfg,
        )
        hit = next(r for r in rows if r["midi"] == 90)
        self.assertEqual(hit["qa_flag"], "REJECTED")
        self.assertFalse(hit["accepted_final"])
        self.assertIsNone(hit["final_value"])

    def test_production_media_matches_combination(self):
        iowa = build_layer("bass_clarinet", "IOWA", "ordinario", "mf", {60: 10.0}, "measured")
        orch = build_layer(
            "bass_clarinet",
            "ORCH",
            "ordinario",
            "mf",
            {60: 30.0},
            "orchidea_family_transfer_estimate",
        )
        only = production_media_of(
            [iowa, orch],
            instrument="bass_clarinet",
            technique="ordinario",
            cfg=CalibrationConfig(combination_method="empirical_only"),
        )
        equal = production_media_of(
            [iowa, orch],
            instrument="bass_clarinet",
            technique="ordinario",
            cfg=CalibrationConfig(combination_method="equal_weight"),
        )
        legacy = production_media_of(
            [iowa, orch],
            instrument="bass_clarinet",
            technique="ordinario",
            cfg=CalibrationConfig(
                combination_method="legacy_equal_weight",
                empirical_weight=0.25,
                transfer_weight=0.75,
            ),
        )
        self.assertEqual(only.cells[60].value, 10.0)
        self.assertEqual(equal.cells[60].value, 20.0)
        self.assertEqual(legacy.cells[60].value, 25.0)
        self.assertEqual(equal.cells[60].origin, "combined_estimate")
        self.assertEqual(only.cells[60].origin, "measured")


class TestCombinedEstimateProvenance(unittest.TestCase):
    def test_blend_is_not_measured_or_high(self):
        cfg = CalibrationConfig(combination_method="equal_weight")
        rows = build_final_rows(
            instrument="bass_clarinet",
            dynamic="mf",
            empirical={60: 10.0},
            transfer={60: 30.0},
            cfg=cfg,
        )
        hit = next(r for r in rows if r["midi"] == 60)
        self.assertEqual(hit["final_value"], 20.0)
        self.assertEqual(hit["provenance"], PROVENANCE_COMBINED_ESTIMATE)
        self.assertNotEqual(hit["provenance"], PROVENANCE_MEASURED)
        self.assertNotEqual(hit["confidence"], QA_HIGH)
        self.assertEqual(hit["qa_flag"], QA_REVIEW)
        self.assertFalse(hit["accepted_final"])

    def test_raw_empirical_row_keeps_measured(self):
        cfg = CalibrationConfig(combination_method="empirical_only")
        rows = build_final_rows(
            instrument="bass_clarinet",
            dynamic="mf",
            empirical={60: 10.0},
            transfer={60: 30.0},
            cfg=cfg,
        )
        hit = next(r for r in rows if r["midi"] == 60)
        self.assertEqual(hit["final_value"], 10.0)
        self.assertEqual(hit["provenance"], PROVENANCE_MEASURED)
        self.assertEqual(hit["confidence"], QA_HIGH)

    def test_mild_blend_uses_established_combined_policy(self):
        cfg = CalibrationConfig(combination_method="equal_weight")
        rows = build_final_rows(
            instrument="bass_clarinet",
            dynamic="mf",
            empirical={60: 10.0},
            transfer={60: 12.0},
            cfg=cfg,
        )
        hit = next(r for r in rows if r["midi"] == 60)
        self.assertEqual(hit["provenance"], PROVENANCE_COMBINED_ESTIMATE)
        self.assertEqual(hit["qa_flag"], QA_MODERATE)
        self.assertTrue(hit["accepted_final"])


class TestConfigurationValidation(unittest.TestCase):
    def test_unsupported_modes_are_rejected(self):
        for name in ("fixed_weight", "validation_optimised", "validation_optimized"):
            cfg = CalibrationConfig(combination_method=name)
            with self.assertRaises(ConfigurationError) as ctx:
                validate_combination_config(cfg)
            self.assertIn("Use empirical_only", str(ctx.exception))

    def test_unknown_mode_is_rejected(self):
        with self.assertRaises(ConfigurationError) as ctx:
            validate_combination_config(CalibrationConfig(combination_method="empircal_only"))
        self.assertIn("Unknown combination method", str(ctx.exception))
        self.assertIn("empirical_only", str(ctx.exception))

    def test_contradictory_transfer_weight_is_rejected(self):
        cfg = CalibrationConfig(
            combination_method="legacy_equal_weight",
            empirical_weight=0.25,
            transfer_weight=0.50,
        )
        with self.assertRaises(ConfigurationError) as ctx:
            validate_combination_config(cfg)
        self.assertIn("w_T = 1 - w_E", str(ctx.exception))

    def test_non_finite_weight_is_rejected(self):
        cfg = CalibrationConfig(empirical_weight=float("nan"), transfer_weight=float("nan"))
        with self.assertRaises(ConfigurationError):
            validate_combination_config(cfg)

    def test_default_yaml_is_unchanged_empirical_only(self):
        cfg = load_config(ROOT / "calibration.yaml")
        self.assertEqual(cfg.combination_method, "empirical_only")
        self.assertEqual(cfg.empirical_weight, 0.5)
        self.assertEqual(cfg.transfer_weight, 0.5)


class TestGuiAndBatchPolicy(unittest.TestCase):
    def test_gui_average_matches_batch_string_policy(self):
        iowa = build_layer("violin", "IOWA", "con sordino", "mf", {60: 10.0}, "measured")
        orch = build_layer("violin", "ORCH", "con sordino", "mf", {60: 30.0}, "measured")
        project = Project(title="synthetic string")
        project.add_layer(iowa)
        project.add_layer(orch)
        batch = production_media_of(
            [iowa, orch], instrument="violin", technique="con sordino"
        )
        gui = production_average_for_layer(project.layers, iowa)
        self.assertEqual(batch.cells[60].value, 20.0)
        self.assertEqual(gui.cells[60].value, batch.cells[60].value)

    def test_gui_average_matches_batch_woodwind_policy(self):
        iowa = build_layer("bass_clarinet", "IOWA", "ordinario", "mf", {60: 10.0}, "measured")
        orch = build_layer(
            "bass_clarinet",
            "ORCH",
            "ordinario",
            "mf",
            {60: 30.0},
            "orchidea_family_transfer_estimate",
        )
        project = Project(title="synthetic woodwind")
        project.add_layer(iowa)
        project.add_layer(orch)
        cfg = CalibrationConfig(combination_method="empirical_only")
        batch = production_media_of(
            [iowa, orch],
            instrument="bass_clarinet",
            technique="ordinario",
            cfg=cfg,
        )
        gui = production_average_for_layer(project.layers, iowa, cfg=cfg)
        self.assertEqual(batch.cells[60].value, 10.0)
        self.assertEqual(gui.cells[60].value, 10.0)

    def test_generic_media_of_remains_distinct_alternative(self):
        iowa = build_layer("bass_clarinet", "IOWA", "ordinario", "mf", {60: 10.0}, "measured")
        orch = build_layer(
            "bass_clarinet",
            "ORCH",
            "ordinario",
            "mf",
            {60: 30.0},
            "orchidea_family_transfer_estimate",
        )
        alternative = media_of([iowa, orch], prefer_measured=False)
        production = production_media_of(
            [iowa, orch],
            instrument="bass_clarinet",
            technique="ordinario",
            cfg=CalibrationConfig(combination_method="empirical_only"),
        )
        self.assertEqual(alternative.cells[60].value, 20.0)
        self.assertEqual(production.cells[60].value, 10.0)

    def test_string_effects_do_not_use_woodwind_priority(self):
        iowa = build_layer("violin", "IOWA", "ordinario", "mf", {60: 10.0}, "measured")
        orch = build_layer("violin", "ORCH", "ordinario", "mf", {60: 30.0}, "measured")
        media = production_media_of([iowa, orch], instrument="violin", technique="ordinario")
        self.assertEqual(media.cells[60].value, 20.0)

    def test_gui_handlers_wire_the_shared_policy(self):
        self.assertIn("production_average_for_layer", inspect.getsource(STEApp._run_media))
        self.assertIn("InsufficientFillData", inspect.getsource(STEApp._run_fill))


class TestFillHandling(unittest.TestCase):
    def test_pchip_two_points_rejected_without_substitution(self):
        layer = build_layer("violin", "IOWA", "ordinario", "mf", {60: 10.0, 62: 12.0}, "measured")
        before = {m: (c.value, c.origin) for m, c in layer.cells.items()}
        with self.assertRaises(InsufficientFillData) as ctx:
            fill_missing(layer, "pchip", max_extrap_semitones=4)
        self.assertIn("at least 3", str(ctx.exception))
        self.assertNotIn("Unknown fill method", str(ctx.exception))
        self.assertEqual({m: (c.value, c.origin) for m, c in layer.cells.items()}, before)

    def test_linear_exterior_is_hold_not_interpolated(self):
        layer = build_layer("violin", "IOWA", "ordinario", "mf", {60: 10.0, 62: 20.0}, "measured")
        layer.range_low = 58
        layer.range_high = 64
        result = fill_missing(layer, "linear", max_extrap_semitones=4)
        self.assertGreater(result.n_extrapolated, 0)
        self.assertEqual(layer.cells[58].origin, "extrapolated_hold")
        self.assertEqual(layer.cells[58].value, 10.0)
        self.assertEqual(layer.cells[61].origin, "interpolated")
        self.assertEqual(layer.cells[64].origin, "extrapolated_hold")
        self.assertEqual(layer.cells[64].value, 20.0)

    def test_pchip_exterior_is_not_ridge(self):
        layer = build_layer(
            "violin",
            "IOWA",
            "ordinario",
            "mf",
            {60: 10.0, 62: 12.0, 64: 11.0},
            "measured",
        )
        layer.range_low = 58
        layer.range_high = 66
        fill_missing(layer, "pchip", max_extrap_semitones=4)
        self.assertEqual(layer.cells[58].origin, "extrapolated_pchip")
        self.assertNotEqual(layer.cells[58].origin, "extrapolated_ridge")
        self.assertIn("extrapolated_pchip", GENERATED_ORIGINS)
        self.assertEqual(normalize_origin("extrapolated_pchip"), "extrapolated_pchip")
        self.assertEqual(normalize_origin("local_ridge_edge"), "extrapolated_ridge")


class TestSyntheticExports(unittest.TestCase):
    def setUp(self):
        remember_transfer_estimates({})

    def tearDown(self):
        remember_transfer_estimates({})

    def _string_project(self) -> Project:
        project = Project(title="synthetic violin")
        for dyn in ("pp", "mf", "ff"):
            project.add_layer(build_layer("violin", "IOWA", "con sordino", dyn, {60: 10.0}, "measured"))
            project.add_layer(build_layer("violin", "ORCH", "con sordino", dyn, {60: 30.0}, "measured"))
        return project

    def _woodwind_project(self) -> Project:
        project = Project(title="synthetic bass clarinet")
        for dyn in ("pp", "mf", "ff"):
            project.add_layer(
                build_layer("bass_clarinet", "IOWA", "ordinario", dyn, {60: 10.0}, "measured")
            )
            project.add_layer(
                build_layer(
                    "bass_clarinet",
                    "ORCH",
                    "ordinario",
                    dyn,
                    {60: 30.0},
                    "orchidea_family_transfer_estimate",
                )
            )
        return project

    def _media_value(self, path: Path, instrument: str, col: int) -> float:
        wb = load_workbook(path, data_only=False)
        ws = wb[media_sheet_name(instrument)]
        note_row = None
        for row in ws.iter_rows(min_row=2, max_col=28, values_only=False):
            if row[27].value == 60:
                note_row = row
                break
        self.assertIsNotNone(note_row)
        cell = note_row[col - 1]
        self.assertFalse(isinstance(cell.value, str) and str(cell.value).startswith("="))
        value = float(cell.value)
        wb.close()
        return value

    def _acoustic_mf(self, path: Path) -> tuple[float, str]:
        wb = load_workbook(path, data_only=False)
        ws = wb["AcousticTable"]
        for rec in ws.iter_rows(min_row=2, values_only=True):
            if rec[2] == 60 and rec[3] == "mf":
                value = rec[4]
                notes = rec[15] or ""
                self.assertFalse(isinstance(value, str) and str(value).startswith("="))
                wb.close()
                return float(value), str(notes)
        wb.close()
        self.fail("AcousticTable missing MIDI 60 mf")

    def test_string_export_media_and_acoustic_are_twenty(self):
        expected = _independent_mean([10.0, 30.0])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "synthetic_violin.xlsx"
            export_zenodo_workbook(self._string_project(), path, "con sordino")
            media_mf = self._media_value(path, "violin", 7)
            media_copy = self._media_value(path, "violin", 14)
            acoustic, _notes = self._acoustic_mf(path)
            self.assertEqual(media_mf, expected)
            self.assertEqual(media_copy, expected)
            self.assertEqual(acoustic, expected)
            wb = load_workbook(path, data_only=False)
            self.assertNotIn("Final_Calibration", wb.sheetnames)
            wb.close()

    def test_woodwind_default_export_keeps_empirical_ten(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "synthetic_bass_clarinet.xlsx"
            export_zenodo_workbook(self._woodwind_project(), path, "ordinario")
            self.assertEqual(self._media_value(path, "bass_clarinet", 7), 10.0)
            acoustic, notes = self._acoustic_mf(path)
            self.assertEqual(acoustic, 10.0)
            wb = load_workbook(path, data_only=False)
            final = wb["Final_Calibration"]
            headers = [cell.value for cell in next(final.iter_rows(min_row=1, max_row=1))]
            midi_i = headers.index("MIDI")
            dyn_i = headers.index("Dynamic")
            val_i = headers.index("Final_Value")
            prov_i = headers.index("Provenance")
            we_i = headers.index("Empirical_Weight")
            cfg_we_i = headers.index("Configured_Empirical_Weight")
            hit = None
            for rec in final.iter_rows(min_row=2, values_only=True):
                if rec[midi_i] == 60 and rec[dyn_i] == "mf":
                    hit = rec
                    break
            self.assertIsNotNone(hit)
            self.assertEqual(hit[val_i], 10.0)
            self.assertEqual(hit[prov_i], PROVENANCE_MEASURED)
            self.assertEqual(hit[we_i], 1.0)
            self.assertEqual(hit[cfg_we_i], 0.5)
            self.assertIn("MEASURED", notes)
            wb.close()

    def test_woodwind_equal_weight_export_is_twenty_combined(self):
        cfg = CalibrationConfig(combination_method="equal_weight")
        with patch("ste_lab.zenodo_export.load_config", return_value=cfg), patch(
            "ste_lab.calibration.load_config", return_value=cfg
        ), patch("ste_lab.media_policy.load_config", return_value=cfg):
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "synthetic_bass_equal.xlsx"
                export_zenodo_workbook(self._woodwind_project(), path, "ordinario")
                self.assertEqual(self._media_value(path, "bass_clarinet", 7), 20.0)
                acoustic, notes = self._acoustic_mf(path)
                self.assertEqual(acoustic, 20.0)
                self.assertIn("COMBINED_ESTIMATE", notes)
                wb = load_workbook(path, data_only=False)
                final = wb["Final_Calibration"]
                headers = [cell.value for cell in next(final.iter_rows(min_row=1, max_row=1))]
                midi_i = headers.index("MIDI")
                dyn_i = headers.index("Dynamic")
                val_i = headers.index("Final_Value")
                prov_i = headers.index("Provenance")
                conf_i = headers.index("Confidence")
                hit = None
                for rec in final.iter_rows(min_row=2, values_only=True):
                    if rec[midi_i] == 60 and rec[dyn_i] == "mf":
                        hit = rec
                        break
                self.assertEqual(hit[val_i], 20.0)
                self.assertEqual(hit[prov_i], PROVENANCE_COMBINED_ESTIMATE)
                self.assertNotEqual(hit[conf_i], QA_HIGH)
                wb.close()

    def test_media_formulas_match_independent_arithmetic(self):
        string_formula = _media_sheet_formula("B", "C", "2", "violin", "con sordino", None)
        self.assertIn("AVERAGE(B2:C2)", string_formula)
        ww = CalibrationConfig(combination_method="empirical_only")
        wood_formula = _media_sheet_formula("B", "C", "2", "bass_clarinet", "ordinario", ww)
        self.assertIn('IF(B2<>"",B2,IF(C2<>"",C2,""))', wood_formula)
        legacy = CalibrationConfig(
            combination_method="legacy_equal_weight",
            empirical_weight=0.25,
            transfer_weight=0.75,
        )
        legacy_formula = _media_sheet_formula("B", "C", "2", "bass_clarinet", "ordinario", legacy)
        self.assertIn("0.25*B2+0.75*C2", legacy_formula)
        self.assertAlmostEqual(0.25 * 10.0 + 0.75 * 30.0, 25.0)


if __name__ == "__main__":
    unittest.main()
