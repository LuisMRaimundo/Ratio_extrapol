# -*- coding: utf-8 -*-
"""Family dispatch, empirical merge, source immutability, string effect exports."""
from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import pandas as pd
from openpyxl import Workbook, load_workbook

from build_ste_preparatory import _folder_is_family_donor, build, find_zenodo_arco
from ste_lab.empirical import assert_not_source, merge_measured_specs, values_equivalent
from ste_lab.notes import midi_to_label
from ste_lab.pipeline import (
    DataDiscoveryError,
    assert_string_effects_survived,
    pipeline_for_instrument,
    run_selected_pipeline,
)
from ste_lab.zenodo_export import _note_grid, measured_ceiling_midi
from run_ste_effects_batch import build_layer


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_compiled(path: Path, notes: dict[int, float]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Research_Core"
    ws.append(["Note", "spectral_mass"])
    for midi, val in notes.items():
        ws.append([midi_to_label(midi), val])
    wb.save(path)


def write_arco_media(path: Path, *, iowa: dict[int, dict[str, float]], orch: dict[int, dict[str, float]], sheet: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["Note", "IOWA pp", "IOWA mf", "IOWA ff", "ORCH pp", "ORCH mf", "ORCH ff"])
    midis = sorted(set(iowa) | set(orch))
    for midi in midis:
        i = iowa.get(midi, {})
        o = orch.get(midi, {})
        ws.append([midi_to_label(midi), i.get("pp"), i.get("mf"), i.get("ff"), o.get("pp"), o.get("mf"), o.get("ff")])
    wb.save(path)


def _stable_compiled(tree: Path, filename: str, notes: dict[int, float]) -> Path:
    path = tree / "_Sustains_Stable" / "analysis_results" / filename
    write_compiled(path, notes)
    return path


def _measured_midis(book: Path, dynamic: str = "pp") -> set[int]:
    wb = load_workbook(book, data_only=True)
    if "Measured_Data" not in wb.sheetnames:
        wb.close()
        return set()
    ws = wb["Measured_Data"]
    headers = [c.value for c in ws[1]]
    out = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        if str(rec.get("collection") or rec.get("Collection") or "").upper() != "IOWA":
            continue
        if str(rec.get("dynamic") or rec.get("Dynamic") or "") != dynamic:
            continue
        midi = rec.get("midi") or rec.get("MIDI")
        if midi is not None:
            out.add(int(midi))
    wb.close()
    return out


class TestDispatch(unittest.TestCase):
    def test_pipeline_for_instrument(self):
        self.assertEqual(pipeline_for_instrument("bass_clarinet"), "woodwinds")
        self.assertEqual(pipeline_for_instrument("viola"), "strings")
        self.assertEqual(pipeline_for_instrument("cello"), "strings")
        self.assertEqual(pipeline_for_instrument("violin"), "strings")
        self.assertEqual(pipeline_for_instrument("double_bass"), "strings")

    def test_selected_runner_is_family_specific(self):
        dummy = Path("unused.xlsx")
        with patch("ste_lab.pipeline.run_woodwind_pipeline", return_value=[]) as woodwind:
            with patch("ste_lab.pipeline.run_string_pipeline", return_value=[]) as strings:
                run_selected_pipeline("bass_clarinet", dummy, dummy)
                woodwind.assert_called_once()
                strings.assert_not_called()
        with patch("ste_lab.pipeline.run_woodwind_pipeline", return_value=[]) as woodwind:
            with patch("ste_lab.pipeline.run_string_pipeline", return_value=[]) as strings:
                run_selected_pipeline("viola", dummy, dummy)
                strings.assert_called_once()
                woodwind.assert_not_called()
        with patch("ste_lab.pipeline.run_woodwind_pipeline", return_value=[]) as woodwind:
            with patch("ste_lab.pipeline.run_string_pipeline", return_value=[]) as strings:
                run_selected_pipeline("cello", dummy, dummy)
                strings.assert_called_once()
                woodwind.assert_not_called()


class TestEmpiricalMerge(unittest.TestCase):
    def test_union_not_longest_wins(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            a = root / "a.xlsx"
            b = root / "b.xlsx"
            write_compiled(a, {48: 10.0, 50: 11.0, 52: 12.0})
            write_compiled(b, {60: 13.0, 62: 14.0})
            specs = [
                {"path": a, "collection": "IOWA", "technique": "ordinario", "dynamic": "pp"},
                {"path": b, "collection": "IOWA", "technique": "ordinario", "dynamic": "pp"},
            ]
            measured, reports = merge_measured_specs(specs)
            midis = set(measured[("IOWA", "ordinario", "pp")]["curve"])
            self.assertEqual(midis, {48, 50, 52, 60, 62})
            self.assertEqual(reports[("IOWA", "ordinario", "pp")].n_dropped, 0)

    def test_conflict_not_averaged(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            a = root / "a.xlsx"
            b = root / "b.xlsx"
            write_compiled(a, {60: 10.0})
            write_compiled(b, {60: 20.0})
            specs = [
                {"path": a, "collection": "IOWA", "technique": "ordinario", "dynamic": "mf"},
                {"path": b, "collection": "IOWA", "technique": "ordinario", "dynamic": "mf"},
            ]
            measured, reports = merge_measured_specs(specs)
            self.assertNotIn(60, measured[("IOWA", "ordinario", "mf")]["curve"])
            self.assertEqual(reports[("IOWA", "ordinario", "mf")].n_conflicts, 1)
            self.assertFalse(values_equivalent(10.0, 20.0))


class TestBassClarinetIngestion(unittest.TestCase):
    def _iowa_tree(self, root: Path, dyn_folder: str, files: list[tuple[str, dict[int, float]]]) -> None:
        tree = root / "BASS_CLARINET_IOWA" / dyn_folder
        for name, notes in files:
            _stable_compiled(tree, name, notes)

    def test_a_multi_file_iowa_merge_reaches_measured_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self._iowa_tree(
                root,
                "pianissimo",
                [
                    ("short_compiled_density_metrics_research.xlsx", {48: 9.0, 50: 9.5}),
                    ("other_compiled_density_metrics_research.xlsx", {60: 10.0, 62: 10.5, 64: 11.0}),
                ],
            )
            out = root / "results"
            written = run_selected_pipeline(
                "bass_clarinet",
                root / "prep.xlsx",
                out,
                rebuild=True,
                rebuild_root=root,
            )
            zen = next(p for p in written if "Zenodo_collections_ordinario" in p.name)
            midis = _measured_midis(zen, "pp")
            self.assertTrue({48, 50, 60, 62, 64}.issubset(midis))

    def test_b_source_media_immutable(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "Bass_clarinet_in_Bb_Zenodo_collections_media.xlsx"
            write_arco_media(
                source,
                iowa={50: {"pp": 8.0, "mf": 9.0, "ff": 10.0}},
                orch={},
                sheet="Bass clarinet in Bb_Media",
            )
            before = _sha256(source)
            self._iowa_tree(root, "pianissimo", [("x_compiled_density_metrics_research.xlsx", {48: 7.0, 50: 8.0})])
            run_selected_pipeline("bass_clarinet", root / "prep.xlsx", root / "results", rebuild=True, rebuild_root=root)
            self.assertEqual(_sha256(source), before)
            generated = list((root / "generated").glob("*generated_media.xlsx"))
            self.assertTrue(generated)
            self.assertNotEqual(generated[0].resolve(), source.resolve())
            with self.assertRaises(RuntimeError):
                assert_not_source(source, {source})

    def test_c_notes_beyond_catalog_sounding_high(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            notes = {37: 12.0, 77: 11.0, 82: 10.0, 84: 9.5}
            self._iowa_tree(root, "mezzo-forte", [("hi_compiled_density_metrics_research.xlsx", notes)])
            written = run_selected_pipeline(
                "bass_clarinet",
                root / "prep.xlsx",
                root / "results",
                rebuild=True,
                rebuild_root=root,
            )
            zen = next(p for p in written if "Zenodo_collections_ordinario" in p.name)
            midis = _measured_midis(zen, "mf")
            self.assertTrue({37, 77, 82, 84}.issubset(midis))
            iowa = build_layer("bass_clarinet", "IOWA", "ordinario", "mf", notes, "measured")
            layers = {("IOWA", "mf"): iowa}
            self.assertEqual(_note_grid(layers, "ordinario", "bass_clarinet")[-1], 84)
            self.assertEqual(measured_ceiling_midi(layers, "bass_clarinet"), 84)
            wb = load_workbook(zen, data_only=True)
            fc = wb["Final_Calibration"]
            headers = [c.value for c in fc[1]]
            for row in fc.iter_rows(min_row=2, values_only=True):
                rec = dict(zip(headers, row))
                if rec.get("Dynamic") == "mf" and rec.get("MIDI") in {82, 84}:
                    self.assertEqual(rec.get("Provenance"), "MEASURED")
            wb.close()


def _string_deposit(root: Path, *, orch_folder: str, media_name: str, media_sheet: str) -> None:
    # Midis sit above every harmonics floor (violin 79, viola 72, cello 60, bass 52).
    iowa = {79: {"pp": 20.0, "mf": 22.0, "ff": 24.0}, 81: {"pp": 21.0, "mf": 23.0, "ff": 25.0}, 83: {"pp": 22.0, "mf": 24.0, "ff": 26.0}}
    orch = {79: {"pp": 10.0, "mf": 11.0, "ff": 12.0}, 81: {"pp": 10.5, "mf": 11.5, "ff": 12.5}, 83: {"pp": 11.0, "mf": 12.0, "ff": 13.0}}
    write_arco_media(root / media_name, iowa=iowa, orch=orch, sheet=media_sheet)
    orch_root = root / orch_folder
    for tech, folder in (
        ("ordinario", "ordinario"),
        ("harmonics", "harmonics"),
        ("sul ponticello", "sul-ponticello"),
        ("con sordino", "con-sord"),
    ):
        vals = {79: 8.0, 81: 8.5, 83: 9.0} if tech != "ordinario" else {79: 11.0, 81: 11.5, 83: 12.0}
        _stable_compiled(
            orch_root / folder / "mezzo-forte",
            f"{folder}_compiled_density_metrics_research.xlsx",
            vals,
        )


def _write_empty_anchors_prep(path: Path) -> None:
    """Preparatory pack with the required sheets but no Anchors_all effects."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Paste_arco"
    ws.append(["note", "midi", "IOWA_pp", "IOWA_mf", "IOWA_ff", "ORCH_pp", "ORCH_mf", "ORCH_ff"])
    ws.append(["C4", 60, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0])
    pe = wb.create_sheet("Paste_effects_mf")
    pe.append(["note", "midi"])
    an = wb.create_sheet("Anchors_all")
    an.append(
        ["effect", "evidence_grade", "dynamic", "note", "midi", "sourceCDM", "targetCDM", "L", "STE_Lab_paste"]
    )
    wb.save(path)


def _expected_string_names(file_stem: str) -> set[str]:
    effects = ("harmonics", "sul_ponticello", "con_sordino", "ordinario")
    out = set()
    for effect in effects:
        out.add(f"{file_stem}_STE_{effect}_IOWA_ORCH.xlsx")
        out.add(f"{file_stem}_Zenodo_collections_{effect}.xlsx")
    return out


class TestStringEffectExports(unittest.TestCase):
    def _run_string(self, instrument: str, orch_folder: str, media_name: str, media_sheet: str, stem: str):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _string_deposit(root, orch_folder=orch_folder, media_name=media_name, media_sheet=media_sheet)
            written = run_selected_pipeline(
                instrument,
                root / "prep.xlsx",
                root / "results",
                rebuild=True,
                rebuild_root=root,
            )
            names = {p.name for p in written}
            self.assertEqual(names, _expected_string_names(stem))
            self.assertFalse(any("tasto" in n.lower() for n in names))
            if stem == "Viola":
                sord = next(p for p in written if p.name.startswith("Viola_STE_con_sordino"))
                wb = load_workbook(sord, data_only=True)
                self.assertNotIn("Final_Calibration", wb.sheetnames)
                self.assertNotIn("Methodology", wb.sheetnames)
                wb.close()
            return written

    def test_d_viola_separate_effect_books(self):
        self._run_string(
            "viola", "ORCH_VLA", "VIOLA_Zenodo_collections_Arco_normal.xlsx", "VIOLA_Media", "Viola"
        )

    def test_e_cello(self):
        self._run_string("cello", "ORCH_CELLO", "Cello_Zenodo_collections_media.xlsx", "Cello_Media", "Cello")

    def test_f_violin(self):
        self._run_string("violin", "Orchidea_violin", "Violin_Zenodo_collections_Arco_normal.xlsx", "Violin_Media", "Violin")

    def test_g_double_bass(self):
        self._run_string(
            "double_bass",
            "ORCH_DBASS",
            "Double_bass_Zenodo_collections_media.xlsx",
            "DBass_Media",
            "Double_bass",
        )

    def test_empty_anchors_with_effect_trees_raises(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_arco_media(
                root / "VIOLA_Zenodo_collections_Arco_normal.xlsx",
                iowa={60: {"pp": 20.0, "mf": 22.0, "ff": 24.0}},
                orch={60: {"pp": 10.0, "mf": 11.0, "ff": 12.0}},
                sheet="VIOLA_Media",
            )
            _stable_compiled(
                root / "ORCH_VLA" / "harmonics" / "mezzo-forte",
                "harm_compiled_density_metrics_research.xlsx",
                {80: 3.0, 81: 3.1, 82: 3.2},
            )
            prep = root / "empty_anchors.xlsx"
            _write_empty_anchors_prep(prep)
            with self.assertRaises(DataDiscoveryError):
                assert_string_effects_survived(root, "viola", prep)


class TestNoCrossFamilyContamination(unittest.TestCase):
    def test_h_woodwind_has_calibration_string_does_not(self):
        self.assertEqual(pipeline_for_instrument("bass_clarinet"), "woodwinds")
        self.assertEqual(pipeline_for_instrument("viola"), "strings")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            tree = root / "BASS_CLARINET_IOWA" / "pianissimo"
            _stable_compiled(tree, "x_compiled_density_metrics_research.xlsx", {50: 8.0, 52: 8.5, 54: 9.0})
            written = run_selected_pipeline(
                "bass_clarinet", root / "p.xlsx", root / "out", rebuild=True, rebuild_root=root
            )
            zen = next(p for p in written if "Zenodo" in p.name)
            wb = load_workbook(zen, data_only=True)
            self.assertIn("Final_Calibration", wb.sheetnames)
            text = " ".join(str(x) for x in wb["Methodology"].iter_rows(values_only=True) for x in x if x)
            self.assertIn("two_log_ratio", text)
            self.assertIn("empirical_only", text)
            wb.close()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _string_deposit(
                root,
                orch_folder="ORCH_VLA",
                media_name="VIOLA_Zenodo_collections_Arco_normal.xlsx",
                media_sheet="VIOLA_Media",
            )
            written = run_selected_pipeline(
                "viola", root / "p.xlsx", root / "out", rebuild=True, rebuild_root=root
            )
            sord = next(p for p in written if "STE_con_sordino" in p.name)
            wb = load_workbook(sord, data_only=True)
            self.assertNotIn("Final_Calibration", wb.sheetnames)
            self.assertNotIn("Methodology", wb.sheetnames)
            wb.close()


def _sheet_headers(path: Path, sheet: str) -> list[str]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    headers = []
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        vals = [str(v).strip() for v in row if v is not None]
        if any(v.lower() == "midi" for v in vals):
            headers = [str(v) if v is not None else "" for v in row]
            break
    wb.close()
    return headers


class TestWoodwindLabelling(unittest.TestCase):
    def test_bass_clarinet_prep_omits_string_headers(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _stable_compiled(
                root / "BASS_CLARINET_IOWA" / "mezzo-forte",
                "x_compiled_density_metrics_research.xlsx",
                {50: 8.0, 60: 9.0, 70: 10.0},
            )
            prep = root / "Bass_clarinet_in_Bb_STE_preparatory.xlsx"
            build(root, "bass_clarinet", prep)
            headers = _sheet_headers(prep, "All_effects_mf")
            joined = " ".join(headers)
            self.assertIn("IOWA_ordinario_mf", headers)
            self.assertIn("ORCH_ordinario_mf", headers)
            self.assertNotIn("IOWA_arco_mf", headers)
            self.assertNotIn("ORCH_arco_mf", headers)
            self.assertNotIn("L_ponticello_mf", headers)
            self.assertNotIn("L_sordino_mf", headers)
            self.assertNotIn("L_harmonics_mf", headers)
            self.assertNotIn("McGill_sordino_mf", headers)
            self.assertNotIn("PHIL_harmonics_mf", headers)
            self.assertNotRegex(joined, r"ponticello|sordino|harmonics|arco")
            eff = _sheet_headers(prep, "Paste_effects_mf")
            self.assertNotIn("McGill_sordino_mf", eff)
            self.assertNotIn("PHIL_harmonics_mf", eff)

    def test_string_prep_keeps_arco_and_effect_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _string_deposit(
                root,
                orch_folder="ORCH_VLA",
                media_name="VIOLA_Zenodo_collections_Arco_normal.xlsx",
                media_sheet="VIOLA_Media",
            )
            prep = root / "Viola_STE_preparatory.xlsx"
            build(root, "viola", prep)
            headers = _sheet_headers(prep, "All_effects_mf")
            self.assertIn("IOWA_arco_mf", headers)
            self.assertIn("ORCH_arco_mf", headers)
            self.assertTrue(any("ponticello" in h or "harmonics" in h or "sordino" in h for h in headers))

    def test_bass_clarinet_measured_data_does_not_stamp_orch_measured(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            notes_i = {50: 10.0, 60: 11.0, 70: 12.0}
            notes_sib = {50: 9.0, 60: 10.0, 70: 11.0}
            notes_orch = {50: 11.0, 60: 12.0, 70: 13.0}
            for dyn in ("pianissimo", "mezzo-forte", "fortissimo"):
                _stable_compiled(
                    root / "BASS_CLARINET_IOWA" / dyn,
                    f"bcl_{dyn}_compiled_density_metrics_research.xlsx",
                    notes_i,
                )
                _stable_compiled(
                    root / "YOWA_Clarinete" / dyn,
                    f"cl_{dyn}_compiled_density_metrics_research.xlsx",
                    notes_sib,
                )
                _stable_compiled(
                    root / "Orchidea_Clarinete" / "ordinario" / dyn,
                    f"orch_{dyn}_compiled_density_metrics_research.xlsx",
                    notes_orch,
                )
            written = run_selected_pipeline(
                "bass_clarinet",
                root / "prep.xlsx",
                root / "results",
                rebuild=True,
                rebuild_root=root,
            )
            zen = next(p for p in written if "Zenodo_collections_ordinario" in p.name)
            wb = load_workbook(zen, data_only=True)
            self.assertIn("Measured_Data", wb.sheetnames)
            ws = wb["Measured_Data"]
            headers = [c.value for c in ws[1]]
            rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
            iowa = [r for r in rows if str(r.get("collection")).upper() == "IOWA"]
            orch = [r for r in rows if str(r.get("collection")).upper() == "ORCH"]
            self.assertTrue(iowa)
            self.assertTrue(all(r.get("provenance") == "MEASURED" for r in iowa))
            self.assertTrue(all(r.get("source_instrument") == "bass_clarinet" for r in iowa))
            self.assertTrue(orch)
            self.assertTrue(all(r.get("provenance") != "MEASURED" for r in orch))
            self.assertTrue(all(r.get("source_instrument") == "clarinet" for r in orch))
            ste = next(p for p in written if p.name.startswith("Bass_clarinet") and "STE" in p.name)
            layer = load_workbook(ste, data_only=True)
            first = [s for s in layer.sheetnames if s not in {"Session_Log", "README"}][0]
            layer_headers = [c.value for c in layer[first][1]]
            self.assertIn("anchor_source", layer_headers)
            self.assertNotIn("anchor_source (arco normal)", layer_headers)
            layer.close()
            wb.close()


class TestEnglishHornFamilyRatio(unittest.TestCase):
    def test_donor_folders(self):
        self.assertTrue(_folder_is_family_donor("OBOE_McGill", "english_horn"))
        self.assertTrue(_folder_is_family_donor("Philharmonia_oboe", "english_horn"))
        self.assertTrue(_folder_is_family_donor("OBOE", "english_horn"))
        self.assertFalse(_folder_is_family_donor("ENGLISH_HORN_McGIll", "english_horn"))
        self.assertFalse(_folder_is_family_donor("Philharmonia_english-horn", "english_horn"))
        self.assertTrue(_folder_is_family_donor("YOWA_Clarinete", "bass_clarinet"))
        self.assertFalse(_folder_is_family_donor("BASS_CLARINET_IOWA", "bass_clarinet"))

    def test_phil_ratio_writes_iowa_and_orch_without_host_iowa(self):
        with tempfile.TemporaryDirectory() as tmp:
            parent = Path(tmp)
            root = parent / "ENGLISH HORN"
            sibling = parent / "OBOE"
            root.mkdir()
            sibling.mkdir()
            _stable_compiled(
                root / "Philharmonia_english-horn" / "mezzo-forte",
                "eh_mf_compiled_density_metrics_research.xlsx",
                {60: 20.0, 61: 22.0},
            )
            _stable_compiled(
                root / "Philharmonia_oboe" / "mezzo-forte",
                "ob_mf_compiled_density_metrics_research.xlsx",
                {60: 10.0, 61: 11.0},
            )
            _stable_compiled(
                sibling / "IOWA_oboe" / "mezzo-forte",
                "iowa_ob_mf_compiled_density_metrics_research.xlsx",
                {60: 8.0, 61: 8.8},
            )
            _stable_compiled(
                sibling / "Orchidea_oboe" / "ordinario" / "mezzo-forte",
                "orch_ob_mf_compiled_density_metrics_research.xlsx",
                {60: 12.0, 61: 13.2},
            )
            prep = root / "English_horn_STE_preparatory.xlsx"
            build(root, "english_horn", prep)
            self.assertTrue(prep.exists())
            gen = root / "generated" / "English_horn_generated_media.xlsx"
            self.assertTrue(gen.exists())
            wb = load_workbook(gen, data_only=True)
            ws = wb["English horn_Media"]
            headers = [c.value for c in ws[1]]
            rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
            by_note = {r["Note"]: r for r in rows}
            self.assertAlmostEqual(by_note["C4"]["IOWA mf"], 16.0, places=6)
            self.assertAlmostEqual(by_note["C4"]["ORCH mf"], 24.0, places=6)
            self.assertAlmostEqual(by_note["Db4"]["IOWA mf"], 17.6, places=6)
            self.assertAlmostEqual(by_note["Db4"]["ORCH mf"], 26.4, places=6)
            prov = {r[0]: r[1] for r in wb["Spine_provenance"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(prov["IOWA"], "family_transfer")
            self.assertNotEqual(prov["ORCH"], "measured")
            wb.close()

    def test_oboe_mcgill_is_not_merged_into_english_horn(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _stable_compiled(
                root / "ENGLISH_HORN_McGIll" / "mf",
                "eh_compiled_density_metrics_research.xlsx",
                {60: 10.0, 62: 11.0},
            )
            _stable_compiled(
                root / "OBOE_McGill" / "mf",
                "ob_compiled_density_metrics_research.xlsx",
                {60: 20.0, 62: 21.0},
            )
            _stable_compiled(
                root / "Philharmonia_english-horn" / "mezzo-forte",
                "eh_phil_compiled_density_metrics_research.xlsx",
                {70: 9.0},
            )
            _stable_compiled(
                root / "Philharmonia_oboe" / "mezzo-forte",
                "ob_phil_compiled_density_metrics_research.xlsx",
                {70: 3.0},
            )
            _stable_compiled(
                root / "IOWA_oboe" / "mezzo-forte",
                "iowa_ob_compiled_density_metrics_research.xlsx",
                {70: 6.0},
            )
            _stable_compiled(
                root / "Orchidea_oboe" / "ordinario" / "mezzo-forte",
                "orch_ob_compiled_density_metrics_research.xlsx",
                {70: 12.0},
            )
            prep = root / "English_horn_STE_preparatory.xlsx"
            build(root, "english_horn", prep)
            wb = load_workbook(prep, data_only=True)
            inv = wb["Measured_inventory"]
            headers = [c.value for c in inv[3]]
            rows = [dict(zip(headers, row)) for row in inv.iter_rows(min_row=4, values_only=True)]
            mcgill = next(
                r
                for r in rows
                if str(r.get("collection")).upper() == "MCGILL"
                and str(r.get("dynamic")) == "mf"
            )
            self.assertEqual(int(mcgill["n"]), 2)
            src = str(mcgill.get("source") or "")
            self.assertIn("eh_compiled", src)
            self.assertNotIn("ob_compiled", src)
            wb.close()


if __name__ == "__main__":
    unittest.main()
