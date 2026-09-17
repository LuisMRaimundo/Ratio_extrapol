# -*- coding: utf-8 -*-
"""Checks that preparatory STE really applies L = ln(target/source).

Run from this folder:

    python -m unittest test_run_ste_from_preparatory.py
"""
from __future__ import annotations

import math
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import pandas as pd
from openpyxl import Workbook

from run_ste_effects_batch import (
    audit_iowa_zenodo_workbook,
    audit_research_tree,
    build_layer,
    load_arco_media,
    transfer_or_copy,
)
from build_ste_preparatory import find_research_trees, find_zenodo_arco
from run_ste_from_preparatory import (
    HARM_LO,
    _instrument_names,
    harm_floor,
    load_preparatory,
    parse_orch_effect_column,
    run,
)
from ste_lab.evidence import evidence_map_rows, principal_empirical_rows, stamp_evidence
from ste_lab.session import Project
from ste_lab.transfer import Anchor, _interp_log_ratio


def _exp_clip(L: float) -> float:
    return math.exp(min(3.0, max(-3.0, L)))


def _write_prep(path: Path) -> dict:
    """Tiny workbook with known numbers. Returns the identities the tests expect."""
    notes = {
        60: "C4",
        67: "G4",
        72: "C5",
        74: "D5",
        79: "G5",
        84: "C6",
    }
    iowa = {
        60: {"pp": 20.0, "mf": 22.0, "ff": 24.0},
        67: {"pp": 21.0, "mf": 23.0, "ff": 25.0},
        72: {"pp": 22.0, "mf": 24.0, "ff": 26.0},
        74: {"pp": 23.0, "mf": 25.0, "ff": 27.0},
        79: {"pp": 24.0, "mf": 26.0, "ff": 28.0},
        84: {"pp": 25.0, "mf": 27.0, "ff": 29.0},
    }
    orch = {
        60: {"pp": 10.0, "mf": 11.0, "ff": 12.0},
        67: {"pp": 10.5, "mf": 11.5, "ff": 12.5},
        72: {"pp": 11.0, "mf": 12.0, "ff": 13.0},
        74: {"pp": 11.5, "mf": 12.5, "ff": 13.5},
        79: {"pp": 12.0, "mf": 13.0, "ff": 14.0},
        84: {"pp": 12.5, "mf": 13.5, "ff": 14.5},
    }
    harm_orch_mf = {60: 99.0, 72: 6.0, 79: 7.0, 84: 8.0}
    pont_orch_mf = {67: 10.35, 72: 10.8, 79: 11.7}
    harm_anchors = [
        (72, 12.0, 6.0),
        (74, 12.5, 9.375),
        (79, 13.0, 7.0),
        (84, 13.5, 8.0),
    ]
    pont_anchors = [
        (67, 11.5, 10.35),
        (72, 12.0, 10.8),
        (79, 13.0, 11.7),
    ]
    sord_anchors = [
        (72, 12.0, 9.6),
        (74, 12.5, 10.0),
        (79, 13.0, 10.4),
    ]

    wb = Workbook()
    arco = wb.active
    arco.title = "Paste_arco"
    arco.append(["Caption — loader must skip this row"])
    arco.append(["note", "midi", "IOWA_pp", "IOWA_mf", "IOWA_ff", "ORCH_pp", "ORCH_mf", "ORCH_ff"])
    for midi, lab in notes.items():
        arco.append(
            [
                lab,
                midi,
                iowa[midi]["pp"],
                iowa[midi]["mf"],
                iowa[midi]["ff"],
                orch[midi]["pp"],
                orch[midi]["mf"],
                orch[midi]["ff"],
            ]
        )

    eff = wb.create_sheet("Paste_effects_mf")
    eff.append(["Caption"])
    eff.append(
        [
            "note",
            "midi",
            "ORCH_harmonics_mf",
            "ORCH_ponticello_mf",
            "PHIL_harmonics_mf",
            "McGill_harmonics_mf",
            "McGill_sordino_mf",
            "McGill_ordinario_mf",
        ]
    )
    for midi, lab in notes.items():
        eff.append(
            [
                lab,
                midi,
                harm_orch_mf.get(midi),
                pont_orch_mf.get(midi),
                5.5 if midi == 72 else None,
                5.0 if midi == 72 else None,
                9.6 if midi in {72, 74, 79} else None,
                12.0 if midi == 72 else None,
            ]
        )

    anc = wb.create_sheet("Anchors_all")
    anc.append(["Caption"])
    anc.append(["effect", "evidence_grade", "note", "midi", "sourceCDM", "targetCDM"])
    for midi, src, tgt in harm_anchors:
        anc.append(["harmonics", "Empirical_ORCH", notes[midi], midi, src, tgt])
    for midi, src, tgt in pont_anchors:
        anc.append(["sul ponticello", "Empirical_ORCH", notes[midi], midi, src, tgt])
    for midi, src, tgt in sord_anchors:
        anc.append(["con sordino", "prediction_McGill", notes[midi], midi, src, tgt])
    for midi in (60, 67, 72, 74):
        anc.append(["sul tasto", "invented", notes[midi], midi, 10.0, 8.0])
    anc.append(["au talon", "Empirical_ORCH", "C5", 72, 10.0, 9.0])
    anc.append(["au talon", "Empirical_ORCH", "D5", 74, 10.0, 9.0])
    wb.save(path)
    return {
        "iowa": iowa,
        "orch": orch,
        "harm_orch_mf": harm_orch_mf,
        "harm_L": {m: math.log(t / s) for m, s, t in harm_anchors},
        "pont_L": {m: math.log(t / s) for m, s, t in pont_anchors},
        "sord_L": {m: math.log(t / s) for m, s, t in sord_anchors},
        "notes": notes,
    }


def _layer_table(path: Path, sheet: str) -> pd.DataFrame:
    with pd.ExcelFile(path) as xl:
        df = pd.read_excel(xl, sheet_name=sheet)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def _value_at(df: pd.DataFrame, midi: int) -> float:
    hit = df.loc[df["MIDI"] == midi]
    if hit.empty:
        raise AssertionError(f"MIDI {midi} missing")
    return float(hit.iloc[0]["Combined density metric"])


def _origin_at(df: pd.DataFrame, midi: int) -> str:
    hit = df.loc[df["MIDI"] == midi]
    if hit.empty:
        raise AssertionError(f"MIDI {midi} missing")
    return str(hit.iloc[0]["Estimate"])


class TestLogRatioTransfer(unittest.TestCase):
    def test_anchor_identity_source_times_exp_L(self):
        anchors = [Anchor(72, 12.0, 6.0), Anchor(79, 13.0, 7.0), Anchor(84, 13.5, 8.0)]
        L = _interp_log_ratio(anchors, [72, 79, 84])
        self.assertAlmostEqual(L[72], math.log(6.0 / 12.0))
        source = build_layer("viola", "IOWA", "ordinario", "mf", {72: 24.0, 79: 26.0, 84: 27.0}, "measured")
        out = transfer_or_copy(source, anchors, "harmonics", "IOWA", "modelled_IOWA_anchored", copy_anchors=False)
        self.assertAlmostEqual(out.cells[72].value, 24.0 * _exp_clip(L[72]))
        self.assertAlmostEqual(out.cells[79].value, 26.0 * _exp_clip(L[79]))
        self.assertNotEqual(out.cells[72].origin.lower(), "measured")

    def test_edge_L_is_held_not_cubically_extrapolated(self):
        anchors = [Anchor(72, 10.0, 5.0), Anchor(78, 10.0, 6.0), Anchor(84, 10.0, 8.0)]
        L = _interp_log_ratio(anchors, [60, 72, 84, 96])
        self.assertAlmostEqual(L[60], math.log(0.5))
        self.assertAlmostEqual(L[96], math.log(0.8))
        self.assertAlmostEqual(L[72], math.log(0.5))
        self.assertAlmostEqual(L[84], math.log(0.8))

    def test_extreme_ratio_is_clipped_at_exp_3(self):
        anchors = [Anchor(72, 1.0, math.exp(4.0)), Anchor(74, 1.0, math.exp(4.0)), Anchor(76, 1.0, math.exp(4.0))]
        source = build_layer("viola", "IOWA", "ordinario", "mf", {72: 2.0}, "measured")
        out = transfer_or_copy(source, anchors, "harmonics", "IOWA", copy_anchors=False)
        self.assertAlmostEqual(out.cells[72].value, 2.0 * math.exp(3.0))

    def test_harmonics_floor_drops_notes_below_c5(self):
        anchors = [Anchor(72, 10.0, 5.0), Anchor(74, 10.0, 5.0), Anchor(79, 10.0, 5.0)]
        source = build_layer(
            "viola", "IOWA", "ordinario", "mf", {67: 20.0, 72: 24.0, 79: 26.0}, "measured"
        )
        out = transfer_or_copy(
            source, anchors, "harmonics", "IOWA", copy_anchors=False, min_midi=HARM_LO
        )
        self.assertNotIn(67, out.cells)
        self.assertIn(72, out.cells)


class TestLoadPreparatory(unittest.TestCase):
    def test_reads_captioned_sheets_and_skips_tasto(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prep.xlsx"
            expect = _write_prep(path)
            pack = load_preparatory(path)
            names = [e["name"] for e in pack["effects"]]
            self.assertIn("harmonics", names)
            self.assertIn("sul ponticello", names)
            self.assertIn("con sordino", names)
            self.assertNotIn("sul tasto", names)
            harm = next(e for e in pack["effects"] if e["name"] == "harmonics")
            self.assertEqual(len(harm["anchors"]), 4)
            self.assertTrue(str(harm["grade"]).startswith("Empirical"))
            self.assertAlmostEqual(pack["arco"]["IOWA"]["mf"][72], expect["iowa"][72]["mf"])
            self.assertAlmostEqual(pack["measured"][("ORCH", "harmonics", "mf")][72][1], 6.0)
            self.assertIn(60, pack["measured"][("ORCH", "harmonics", "mf")])

    def test_missing_sheet_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.xlsx"
            wb = Workbook()
            wb.active.title = "README"
            wb.save(path)
            with self.assertRaises(SystemExit):
                load_preparatory(path)

    def test_instrument_slug(self):
        self.assertEqual(_instrument_names("Viola"), ("viola", "Viola"))
        self.assertEqual(_instrument_names("double bass"), ("double_bass", "Double_bass"))


class TestRunFromPreparatory(unittest.TestCase):
    def test_exported_books_obey_L_and_evidence_rules(self):
        with tempfile.TemporaryDirectory() as tmp:
            prep = Path(tmp) / "prep.xlsx"
            out = Path(tmp) / "out"
            expect = _write_prep(prep)
            written = run(prep, out, instrument="viola")
            names = {p.name for p in written}
            self.assertEqual(
                names,
                {
                    "Viola_STE_harmonics_IOWA_ORCH.xlsx",
                    "Viola_Zenodo_collections_harmonics.xlsx",
                    "Viola_STE_sul_ponticello_IOWA_ORCH.xlsx",
                    "Viola_Zenodo_collections_sul_ponticello.xlsx",
                    "Viola_STE_con_sordino_IOWA_ORCH.xlsx",
                    "Viola_Zenodo_collections_con_sordino.xlsx",
                    "Viola_STE_ordinario_IOWA_ORCH.xlsx",
                    "Viola_Zenodo_collections_ordinario.xlsx",
                },
            )
            self.assertFalse(any("tasto" in n.lower() for n in names))
            self.assertFalse(any("talon" in n.lower() for n in names))

            harm = out / "Viola_STE_harmonics_IOWA_ORCH.xlsx"
            iowa_mf = _layer_table(harm, "viola_IOWA_mf")
            orch_mf = _layer_table(harm, "viola_ORCH_mf")
            iowa_pp = _layer_table(harm, "viola_IOWA_pp")

            L72 = expect["harm_L"][72]
            self.assertAlmostEqual(
                _value_at(iowa_mf, 72),
                expect["iowa"][72]["mf"] * _exp_clip(L72),
            )
            self.assertAlmostEqual(
                _value_at(iowa_pp, 72),
                expect["iowa"][72]["pp"] * _exp_clip(L72),
            )
            self.assertAlmostEqual(_value_at(orch_mf, 72), expect["harm_orch_mf"][72])
            self.assertEqual(_origin_at(orch_mf, 72).lower(), "measured")
            self.assertNotIn(74, set(orch_mf["MIDI"]))
            self.assertTrue((iowa_mf["MIDI"] >= HARM_LO).all())
            self.assertTrue((orch_mf["MIDI"] >= HARM_LO).all())
            self.assertNotIn(60, set(orch_mf["MIDI"]))

            media_72 = 0.5 * (_value_at(iowa_mf, 72) + _value_at(orch_mf, 72))
            media = _layer_table(harm, "Media_mf")
            self.assertAlmostEqual(_value_at(media, 72), media_72)

            pont = out / "Viola_STE_sul_ponticello_IOWA_ORCH.xlsx"
            pont_iowa = _layer_table(pont, "viola_IOWA_mf")
            self.assertAlmostEqual(
                _value_at(pont_iowa, 72),
                expect["iowa"][72]["mf"] * _exp_clip(expect["pont_L"][72]),
            )

            sord = out / "Viola_STE_con_sordino_IOWA_ORCH.xlsx"
            sord_iowa = _layer_table(sord, "viola_IOWA_mf")
            sord_orch = _layer_table(sord, "viola_ORCH_mf")
            self.assertAlmostEqual(
                _value_at(sord_iowa, 72),
                expect["iowa"][72]["mf"] * _exp_clip(expect["sord_L"][72]),
            )
            self.assertAlmostEqual(
                _value_at(sord_orch, 72),
                expect["orch"][72]["mf"] * _exp_clip(expect["sord_L"][72]),
            )
            self.assertNotEqual(_origin_at(sord_orch, 72).lower(), "measured")
            mcgill = _layer_table(sord, "viola_MCGILL_mf")
            self.assertAlmostEqual(_value_at(mcgill, 72), 9.6)
            em = pd.read_excel(sord, sheet_name="Evidence_Map")
            mcgill_map = em[em["collection"].astype(str).str.upper() == "MCGILL"]
            self.assertFalse(mcgill_map.empty)
            self.assertEqual(str(mcgill_map.iloc[0]["evidence_role"]), "context")
            self.assertEqual(str(mcgill_map.iloc[0]["principal_evidence"]), "no")
            self.assertNotIn("orchidea", str(mcgill_map.iloc[0]["note"]).lower())

            with pd.ExcelFile(sord) as xl:
                readme = pd.read_excel(xl, sheet_name="README")
            notes = " ".join(str(v) for v in readme.iloc[:, 1].tolist()).lower()
            self.assertIn("prediction", notes)


class TestEvidenceStamps(unittest.TestCase):
    def test_iowa_same_dynamic_is_modelled_not_inherited(self):
        layer = build_layer("cello", "IOWA", "con sordino", "ff", {72: 10.0}, "modelled_IOWA_anchored")
        stamp_evidence(layer, inherited_from=None, prediction=False)
        self.assertEqual(layer.labels["evidence_role"], "modelled_completion")
        self.assertNotEqual((layer.labels.get("l_donor_dynamic") or "").lower(), "mf")

    def test_iowa_unrecorded_dynamic_is_inherited(self):
        layer = build_layer("cello", "IOWA", "harmonics", "pp", {72: 10.0}, "modelled_IOWA_anchored")
        stamp_evidence(layer, inherited_from="mf", prediction=False)
        self.assertEqual(layer.labels["evidence_role"], "inherited_dynamic")
        self.assertEqual(layer.labels["l_donor_dynamic"], "mf")

    def test_mcgill_and_phil_are_context_not_principal(self):
        mcgill = build_layer("cello", "MCGILL", "con sordino", "mf", {72: 9.6}, "measured")
        phil = build_layer("cello", "PHIL", "harmonics", "mf", {72: 4.0}, "measured")
        stamp_evidence(mcgill)
        stamp_evidence(phil)
        project = Project(title="stamp check")
        project.add_layer(mcgill)
        project.add_layer(phil)
        sord_rows = evidence_map_rows(project, "con sordino")
        self.assertEqual(sord_rows[0]["evidence_role"], "context")
        self.assertEqual(sord_rows[0]["principal_evidence"], "no")
        self.assertNotIn("Orchidea", sord_rows[0]["note"])
        harm_rows = evidence_map_rows(project, "harmonics")
        self.assertEqual(harm_rows[0]["evidence_role"], "context")
        self.assertEqual(harm_rows[0]["principal_evidence"], "no")

    def test_orch_recorded_is_principal(self):
        layer = build_layer("cello", "ORCH", "con sordino", "ff", {72: 8.0}, "measured")
        stamp_evidence(layer)
        project = Project(title="stamp check")
        project.add_layer(layer)
        row = evidence_map_rows(project, "con sordino")[0]
        self.assertEqual(row["evidence_role"], "empirical")
        self.assertEqual(row["principal_evidence"], "yes")


class TestWoodwindPrincipalEvidence(unittest.TestCase):
    def test_iowa_ordinario_is_principal_on_woodwind(self):
        layer = build_layer("bass_clarinet", "IOWA", "ordinario", "mf", {37: 10.0, 48: 12.0}, "measured")
        stamp_evidence(layer)
        project = Project(title="ww")
        project.add_layer(layer)
        row = evidence_map_rows(project, "ordinario")[0]
        self.assertEqual(row["evidence_role"], "empirical")
        self.assertEqual(row["principal_evidence"], "yes")
        rows = principal_empirical_rows(project, "ordinario")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["collection"], "IOWA")
        self.assertEqual(rows[0]["value"], 10.0)

    def test_family_transfer_orch_is_not_principal(self):
        iowa = build_layer("bass_clarinet", "IOWA", "ordinario", "mf", {48: 12.0}, "measured")
        orch = build_layer("bass_clarinet", "ORCH", "ordinario", "mf", {48: 11.0}, "family_transfer")
        stamp_evidence(iowa)
        stamp_evidence(orch)
        project = Project(title="ww")
        project.add_layer(iowa)
        project.add_layer(orch)
        rows = principal_empirical_rows(project, "ordinario")
        self.assertEqual([r["collection"] for r in rows], ["IOWA"])
        em = evidence_map_rows(project, "ordinario")
        orch_row = next(r for r in em if str(r["collection"]).upper() == "ORCH")
        self.assertEqual(orch_row["principal_evidence"], "no")

    def test_string_iowa_ordinario_does_not_fill_empty_orch_sheet(self):
        iowa = build_layer("cello", "IOWA", "ordinario", "mf", {48: 12.0}, "measured")
        stamp_evidence(iowa)
        project = Project(title="str")
        project.add_layer(iowa)
        self.assertEqual(principal_empirical_rows(project, "ordinario"), [])
        self.assertEqual(evidence_map_rows(project, "ordinario")[0]["principal_evidence"], "no")


class TestZenodoSheetRefs(unittest.TestCase):
    def test_sheet_ref_quotes_spaces(self):
        from ste_lab.zenodo_export import _sheet_ref

        self.assertEqual(_sheet_ref("Violin_IOWA_pp"), "Violin_IOWA_pp")
        self.assertEqual(
            _sheet_ref("Bass clarinet in Bb_IOWA_pp"),
            "'Bass clarinet in Bb_IOWA_pp'",
        )

    def test_media_formulas_quote_bass_clarinet_sheets(self):
        from openpyxl import Workbook
        from ste_lab.zenodo_export import _write_media

        wb = Workbook()
        ws = wb.active
        _write_media(ws, [37], 77, "bass_clarinet")
        formula = str(ws.cell(2, 2).value)
        self.assertIn("'Bass clarinet in Bb_IOWA_pp'", formula)
        self.assertNotIn("=IF(Bass clarinet", formula)

    def test_media_copies_layer_values(self):
        from openpyxl import Workbook
        from ste_lab.zenodo_export import _write_media

        layers = {
            ("IOWA", "pp"): build_layer("bass_clarinet", "IOWA", "ordinario", "pp", {37: 30.0}, "measured"),
            ("ORCH", "pp"): build_layer("bass_clarinet", "ORCH", "ordinario", "pp", {37: 20.0}, "family_transfer"),
            ("IOWA", "mf"): build_layer("bass_clarinet", "IOWA", "ordinario", "mf", {37: 40.0}, "measured"),
            ("ORCH", "mf"): build_layer("bass_clarinet", "ORCH", "ordinario", "mf", {37: 22.0}, "family_transfer"),
        }
        wb = Workbook()
        ws = wb.active
        _write_media(ws, [37], 77, "bass_clarinet", layers=layers)
        self.assertEqual(ws.cell(2, 1).value, "Db2")
        self.assertAlmostEqual(ws.cell(2, 2).value, 30.0)
        self.assertAlmostEqual(ws.cell(2, 3).value, 20.0)
        self.assertAlmostEqual(ws.cell(2, 4).value, 30.0)
        self.assertAlmostEqual(ws.cell(2, 5).value, 40.0)
        self.assertAlmostEqual(ws.cell(2, 6).value, 22.0)
        self.assertAlmostEqual(ws.cell(2, 7).value, 40.0)


class TestOrchideaRecordedRoutine(unittest.TestCase):
    def test_column_parser(self):
        self.assertEqual(parse_orch_effect_column("ORCH_sordino_ff"), ("con sordino", "ff"))
        self.assertEqual(parse_orch_effect_column("ORCH_ponticello_pp"), ("sul ponticello", "pp"))
        self.assertEqual(parse_orch_effect_column("ORCH_harmonics_mf"), ("harmonics", "mf"))
        self.assertIsNone(parse_orch_effect_column("ORCH_arco_mf"))
        self.assertIsNone(parse_orch_effect_column("ORCH_tasto_mf"))
        self.assertIsNone(parse_orch_effect_column("McGill_sordino_mf"))

    def test_recorded_orch_dynamic_is_not_replaced_by_L(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "prep.xlsx"
            out = Path(tmp) / "out"
            wb = Workbook()
            arco = wb.active
            arco.title = "Paste_arco"
            arco.append(["note", "midi", "IOWA_pp", "IOWA_mf", "IOWA_ff", "ORCH_pp", "ORCH_mf", "ORCH_ff"])
            for midi, lab, iowa, orch in ((72, "C5", 24.0, 12.0), (79, "G5", 26.0, 13.0), (84, "C6", 28.0, 14.0)):
                arco.append([lab, midi, iowa - 2, iowa, iowa + 2, orch - 1, orch, orch + 1])
            eff = wb.create_sheet("Paste_effects_mf")
            eff.append(["note", "midi", "ORCH_sordino_pp", "ORCH_sordino_mf", "ORCH_sordino_ff"])
            for midi, lab, pp, mf, ff in (
                (72, "C5", 9.0, 9.6, 8.0),
                (79, "G5", 10.0, 10.4, 9.0),
                (84, "C6", 11.0, 11.2, 10.0),
            ):
                eff.append([lab, midi, pp, mf, ff])
            anc = wb.create_sheet("Anchors_all")
            anc.append(["effect", "evidence_grade", "dynamic", "note", "midi", "sourceCDM", "targetCDM"])
            for midi, src, tgt, dyn in (
                (72, 12.0, 9.6, "mf"),
                (79, 13.0, 10.4, "mf"),
                (84, 14.0, 11.2, "mf"),
                (72, 13.0, 8.0, "ff"),
                (79, 14.0, 9.0, "ff"),
                (84, 15.0, 10.0, "ff"),
            ):
                lab = {72: "C5", 79: "G5", 84: "C6"}[midi]
                anc.append(["con sordino", "Empirical_ORCH", dyn, lab, midi, src, tgt])
            wb.save(path)
            written = run(path, out, instrument="cello")
            sord = out / "Cello_STE_con_sordino_IOWA_ORCH.xlsx"
            self.assertTrue(sord.exists(), written)
            orch_ff = _layer_table(sord, "cello_ORCH_ff")
            self.assertAlmostEqual(_value_at(orch_ff, 72), 8.0)
            self.assertEqual(_origin_at(orch_ff, 72).lower(), "measured")
            orch_pp = _layer_table(sord, "cello_ORCH_pp")
            self.assertAlmostEqual(_value_at(orch_pp, 72), 9.0)
            self.assertEqual(_origin_at(orch_pp, 72).lower(), "measured")
            iowa_ff = _layer_table(sord, "cello_IOWA_ff")
            self.assertAlmostEqual(_value_at(iowa_ff, 72), 26.0 * _exp_clip(math.log(8.0 / 13.0)))
            em = pd.read_excel(sord, sheet_name="Evidence_Map")
            def _row(coll, dyn):
                hit = em[(em["collection"].astype(str).str.upper() == coll) & (em["dynamic"].astype(str) == dyn)]
                self.assertFalse(hit.empty, f"{coll} {dyn} missing from Evidence_Map")
                return hit.iloc[0]
            orch_ff_map = _row("ORCH", "ff")
            self.assertEqual(str(orch_ff_map["evidence_role"]), "empirical")
            self.assertEqual(str(orch_ff_map["principal_evidence"]), "yes")
            iowa_ff_map = _row("IOWA", "ff")
            self.assertEqual(str(iowa_ff_map["evidence_role"]), "modelled_completion")
            self.assertNotEqual(str(iowa_ff_map["l_donor_dynamic"]).strip().lower(), "mf")
            iowa_pp_map = _row("IOWA", "pp")
            self.assertEqual(str(iowa_pp_map["evidence_role"]), "inherited_dynamic")
            self.assertEqual(str(iowa_pp_map["l_donor_dynamic"]).strip().lower(), "mf")


class TestResearchTreeWarnings(unittest.TestCase):
    def test_missing_folder_warns(self):
        missing = Path(tempfile.gettempdir()) / "ste_lab_no_such_research_tree"
        warns = audit_research_tree(missing, technique="harmonics")
        self.assertTrue(any("missing research folder" in w for w in warns))

    def test_empty_tree_warns_stable_and_compiled(self):
        with tempfile.TemporaryDirectory() as tmp:
            tree = Path(tmp) / "harmonics"
            tree.mkdir()
            warns = "\n".join(audit_research_tree(tree, technique="harmonics"))
            self.assertIn("no _Sustains_Stable", warns)
            self.assertIn("compiled_density_metrics_research.xlsx", warns)

    def test_stable_without_compiled_warns(self):
        with tempfile.TemporaryDirectory() as tmp:
            tree = Path(tmp) / "harmonics"
            (tree / "mezzo-forte" / "_Sustains_Stable").mkdir(parents=True)
            warns = "\n".join(audit_research_tree(tree, technique="harmonics"))
            self.assertIn("_Sustains_Stable has no", warns)

    def test_stable_with_prefixed_compiled_is_quiet(self):
        with tempfile.TemporaryDirectory() as tmp:
            tree = Path(tmp) / "harmonics"
            stable = tree / "mezzo-forte" / "_Sustains_Stable"
            stable.mkdir(parents=True)
            (stable / "viola_harmonics_mezzo-forte_compiled_density_metrics_research.xlsx").write_bytes(b"PK")
            self.assertEqual(audit_research_tree(tree, technique="harmonics"), [])


class TestIowaZenodoWorkbook(unittest.TestCase):
    def test_missing_iowa_book_warns(self):
        missing = Path(tempfile.gettempdir()) / "no_such_iowa_zenodo.xlsx"
        warns = audit_iowa_zenodo_workbook(missing)
        self.assertTrue(any("missing IOWA Zenodo" in w for w in warns))

    def test_book_without_iowa_sheets_warns(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Empty_Zenodo.xlsx"
            wb = Workbook()
            wb.active.title = "README"
            wb.save(path)
            warns = "\n".join(audit_iowa_zenodo_workbook(path))
            self.assertIn("no Media sheet", warns)
            self.assertIn("by-string sheet", warns)

    def test_by_string_sheet_fills_iowa_when_media_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Viola_Zenodo.xlsx"
            wb = Workbook()
            media = wb.active
            media.title = "VIOLA_Media"
            media.append(["Note", "ORCH mf"])
            media.append(["C5", 12.0])
            iowa = wb.create_sheet("VIOLA__IOWA_mf")
            iowa.append(["Source note", "String", "CDM", "CDM - Media"])
            iowa.append(["C5", "C STRING", 24.0, 24.0])
            wb.create_sheet("VIOLA_IOWA_pp")
            wb.create_sheet("VIOLA_IOWA_ff")
            wb.save(path)
            curves = load_arco_media(path, media_sheet="VIOLA_Media")
            self.assertAlmostEqual(curves["IOWA"]["mf"][72], 24.0)
            self.assertAlmostEqual(curves["ORCH"]["mf"][72], 12.0)


class TestRebuildFolderHunt(unittest.TestCase):
    def test_finds_arco_not_ste_or_effect_books(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "CELLO_Zenodo_collections_Arco_normal.xlsx").write_bytes(b"PK")
            (root / "Cello_STE_harmonics_IOWA_ORCH.xlsx").write_bytes(b"PK")
            (root / "Cello_Zenodo_collections_harmonics.xlsx").write_bytes(b"PK")
            (root / "Cello_Zenodo_collections_con_sordino.xlsx").write_bytes(b"PK")
            got = find_zenodo_arco(root)
            self.assertIsNotNone(got)
            self.assertEqual(got.name, "CELLO_Zenodo_collections_Arco_normal.xlsx")

    def test_trees_skip_tasto_and_iowa(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "Orchidea_cello" / "harmonics").mkdir(parents=True)
            (root / "Orchidea_cello" / "con-sord").mkdir()
            (root / "Orchidea_cello" / "sul-tasto").mkdir()
            (root / "Philharmonia_cello").mkdir()
            (root / "McGill_cello").mkdir()
            (root / "IOWA_CELLO").mkdir()
            trees = find_research_trees(root)
            labels = {(coll, tech or "", path.name) for path, coll, tech in trees}
            self.assertIn(("ORCH", "harmonics", "harmonics"), labels)
            self.assertIn(("ORCH", "con sordino", "con-sord"), labels)
            self.assertIn(("PHIL", "", "Philharmonia_cello"), labels)
            self.assertIn(("MCGILL", "", "McGill_cello"), labels)
            self.assertFalse(any("tasto" in path.name.lower() or tech == "sul tasto" for path, _, tech in trees))
            self.assertFalse(any(coll == "IOWA" or "iowa" in path.name.lower() for path, coll, _ in trees))

    def test_harm_floor_by_instrument(self):
        self.assertEqual(harm_floor("viola"), 72)
        self.assertEqual(harm_floor("cello"), 60)
        self.assertEqual(harm_floor("violin"), 79)
        self.assertEqual(harm_floor("double bass"), 52)

    def test_bass_clarinet_skips_sibling_clarinet_media(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "Clarinet_Zenodo_collections_media.xlsx").write_bytes(b"PK")
            (root / "Bass_clarinet_in_Bb_Zenodo_collections_media.xlsx").write_bytes(b"PK")
            got = find_zenodo_arco(root, "bass_clarinet")
            self.assertIsNotNone(got)
            self.assertEqual(got.name, "Bass_clarinet_in_Bb_Zenodo_collections_media.xlsx")

    def test_bass_clarinet_trees_keep_iowa_drop_sibling_clarinet(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "BASS_CLARINET_IOWA").mkdir()
            (root / "Philharmonia_bass clarinet").mkdir()
            (root / "Philharmonia_clarinet").mkdir()
            (root / "Orchidea_Clarinete" / "ordinario").mkdir(parents=True)
            (root / "McGill_CLARINET_BASS").mkdir()
            trees = find_research_trees(root, "bass_clarinet")
            names = {path.name for path, _, _ in trees}
            self.assertIn("BASS_CLARINET_IOWA", names)
            self.assertIn("Philharmonia_bass clarinet", names)
            self.assertIn("McGill_CLARINET_BASS", names)
            self.assertNotIn("Philharmonia_clarinet", names)
            self.assertFalse(any("Orchidea" in path.name or "Clarinete" in path.name for path, _, _ in trees))


if __name__ == "__main__":
    unittest.main(verbosity=2)
