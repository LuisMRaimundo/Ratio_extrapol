# -*- coding: utf-8 -*-
"""Rebuild one STE preparatory workbook from a deposit folder you choose."""
from __future__ import annotations

import argparse
import math
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from run_ste_effects_batch import (
    DYN_MAP,
    audit_iowa_zenodo_workbook,
    audit_research_tree,
    load_arco_media,
    load_spectral_mass,
    log_discovered_specs,
)
from ste_lab.catalog import family_donor_instrument, orchestral_group, resolve_instrument
from ste_lab.notes import midi_to_label

CORE = ("pp", "mf", "ff")
HARM_FLOOR = {
    "violin": 79,
    "viola": 72,
    "cello": 60,
    "double_bass": 52,
}
DEPOSIT_DEFAULTS: dict[str, dict] = {
    "viola": {
        "root": Path(r"D:\CORDAS_3\VIOLA 4"),
        "prep_name": "Viola_STE_preparatory.xlsx",
        "out_name": "Viola_Extrapoled_effects",
        "media_sheet": "VIOLA_Media",
    },
    "cello": {
        "root": Path(r"D:\CORDAS_3\CELLO"),
        "prep_name": "Cello_STE_preparatory.xlsx",
        "out_name": "Cello_Extrapoled_effects",
        "media_sheet": "Cello_Media",
    },
    "violin": {
        "root": Path(r"D:\CORDAS_3\VIOLIN 4"),
        "prep_name": "Violin_STE_preparatory.xlsx",
        "out_name": "Extrapoled_effects",
        "media_sheet": "Violin_Media",
    },
    "double_bass": {
        "root": Path(r"D:\CORDAS_3\DOUBLE_BASS"),
        "prep_name": "Double_bass_STE_preparatory.xlsx",
        "out_name": "Double_bass_Extrapoled_effects",
        "media_sheet": "DBass_Media",
    },
    "english_horn": {
        "root": Path(r"D:\MADEIRAS_2\ENGLISH HORN"),
        "prep_name": "English_horn_STE_preparatory.xlsx",
        "out_name": "English_horn results",
        "media_sheet": "English horn_Media",
    },
    "bass_clarinet": {
        "root": Path(r"D:\MADEIRAS_2\BASS_CLARINET"),
        "prep_name": "Bass_clarinet_in_Bb_STE_preparatory.xlsx",
        "out_name": "Bass_clarinet results",
        "media_sheet": "Bass clarinet in Bb_Media",
    },
}
ORCH_TECH_SUBS = {
    "sul-ponticello": "sul ponticello",
    "arco-sul-ponticello": "sul ponticello",
    "sul-tasto": "sul tasto",
    "arco-sul-tasto": "sul tasto",
    "harmonics": "harmonics",
    "ordinario": "ordinario",
    "arco-normal": "ordinario",
    "con-sord": "con sordino",
    "con-sordino": "con sordino",
    "muted": "con sordino",
}

last_path: Path | None = None
last_warnings: list[str] = []


def deposit_defaults(instrument: str) -> dict | None:
    spec = resolve_instrument(instrument)
    key = spec.instrument_id if spec else (instrument or "").strip().lower().replace(" ", "_")
    return DEPOSIT_DEFAULTS.get(key)


def known_roots() -> set[Path]:
    return {d["root"] for d in DEPOSIT_DEFAULTS.values()}


def known_prep_paths() -> set[Path]:
    return {d["root"] / d["prep_name"] for d in DEPOSIT_DEFAULTS.values()}


def known_out_paths() -> set[Path]:
    outs = {d["root"] / d["out_name"] for d in DEPOSIT_DEFAULTS.values()}
    outs.update(known_roots())
    return outs


def harm_floor(instrument: str) -> int | None:
    spec = resolve_instrument(instrument)
    key = spec.instrument_id if spec else (instrument or "").strip().lower().replace(" ", "_")
    return HARM_FLOOR.get(key)


def display_name(instrument: str) -> str:
    spec = resolve_instrument(instrument)
    if spec:
        return spec.display_name
    return (instrument or "instrument").replace("_", " ").title()


def find_zenodo_arco(root: Path, instrument: str | None = None) -> Path | None:
    """Zenodo collections arco/media book. Never a STE result book."""
    cands: list[Path] = []
    if not root.exists():
        return None
    spec = resolve_instrument(instrument or "")
    instr_id = spec.instrument_id if spec else ""
    for path in root.glob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        name = path.name.lower()
        if "zenodo_collections" not in name:
            continue
        if "ste_" in name:
            continue
        if "_generated_media" in name or "generated" in path.parts:
            continue
        if any(
            token in name
            for token in ("harmonics", "ponticello", "sordino", "sul_tasto", "sul-tasto", "sulta")
        ):
            continue
        if "arco" in name or "collections_media" in name or name.endswith("_media.xlsx"):
            if instr_id == "bass_clarinet" and "clarinet" in name and "bass" not in name:
                continue
            cands.append(path)
    if not cands:
        return None
    cands.sort(key=lambda p: (0 if "arco" in p.name.lower() else 1, p.name.lower()))
    return cands[0]


def _fold_folder_name(name: str) -> str:
    n = name.lower().replace(" ", "_").replace("-", "_")
    return n.translate(
        str.maketrans("áàâãäéèêëíìîïóòôõöúùûüç", "aaaaaeeeeiiiiooooouuuuc")
    )


def _folder_is_family_donor(name: str, instrument: str | None) -> bool:
    """True when the folder is a sibling in the family, not the target instrument."""
    spec = resolve_instrument(instrument or "")
    if not spec:
        return False
    n = name.lower().replace(" ", "_").replace("-", "_")
    if spec.instrument_id == "bass_clarinet":
        target = (
            "bass_clarinet",
            "bassclarinet",
            "clarinet_bass",
            "clarinete_baixo",
            "clar_baixo",
            "bass_clar",
        )
        if any(tok in n for tok in target):
            return False
        return "clarinet" in n or "clarinete" in n or n.startswith("yowa")
    donor = family_donor_instrument(spec.instrument_id)
    if not donor:
        return False
    n = _fold_folder_name(name)
    host = (
        spec.instrument_id,
        spec.instrument_id.replace("_", ""),
        spec.display_name.lower().replace(" ", "_"),
        spec.display_name.lower().replace(" ", ""),
    )
    if spec.instrument_id == "english_horn":
        host += ("cor_anglais", "coranglais", "corne_ingles")
    if spec.instrument_id == "contrabassoon":
        host += ("contra_bassoon", "contrabasson")
    if any(tok in n for tok in host if tok):
        return False
    donor_needles = {
        "oboe": ("oboe",),
        "flute": ("flute", "flauta"),
        "bassoon": ("bassoon", "basson", "fagot", "fagott"),
        "clarinet": ("clarinet", "clarinete", "yowa"),
    }.get(donor, (donor,))
    return any(tok in n for tok in donor_needles)


def _sibling_family_deposits(root: Path, instrument: str | None) -> list[Path]:
    """Neighbouring deposit that holds the family donor's Iowa / Orchidea trees."""
    spec = resolve_instrument(instrument or "")
    if not spec or spec.instrument_id == "bass_clarinet":
        return []
    if not family_donor_instrument(spec.instrument_id) or not root.exists():
        return []
    skip = ("full analysis", "full_analysis", "results", "generated", "extrapoled", "extrapolated")
    found: list[Path] = []
    for child in sorted(root.parent.iterdir(), key=lambda p: p.name.lower()):
        if not child.is_dir() or child.resolve() == root.resolve():
            continue
        nl = child.name.lower()
        if any(tok in nl for tok in skip):
            continue
        if _folder_is_family_donor(child.name, instrument):
            found.append(child)
    return found


def _dyn_from_parts(parts: list[str]) -> str | None:
    for folder, code in DYN_MAP.items():
        if folder in parts:
            return code
    for part in parts:
        for code in ("pp", "mp", "mf", "ff"):
            if part == code or part.endswith("_" + code):
                return code
    for part in parts:
        for code in ("p", "f"):
            if part == code or part.endswith("_" + code):
                return code
    return None


def find_research_trees(
    root: Path, instrument: str | None = None
) -> list[tuple[Path, str, str | None]]:
    """Target-instrument Orchidea / Philharmonia / McGill / (woodwind) Iowa trees.

    String deposits still skip Iowa folders (Iowa lives on the Zenodo book).
    Woodwind deposits include Iowa compiled trees for the *target* instrument.
    Sibling family folders (e.g. Bb clarinet inside a bass-clarinet deposit) are skipped.
    """
    trees: list[tuple[Path, str, str | None]] = []
    if not root.exists():
        return trees
    woodwind = orchestral_group(instrument or "") == "woodwinds"
    for child in sorted(root.iterdir(), key=lambda p: p.name.lower()):
        if not child.is_dir():
            continue
        name = child.name.lower()
        if name.startswith("."):
            continue
        if "extrapoled" in name or "extrapolated" in name or "results" in name:
            continue
        if "tasto" in name:
            continue
        if _folder_is_family_donor(child.name, instrument):
            continue
        iowa_name = name.startswith("iowa") or "iowa_" in name or name.endswith("_iowa")
        if iowa_name and not woodwind:
            continue
        if "orch" in name or "orchidea" in name:
            found_tech = False
            for sub in sorted(child.iterdir(), key=lambda p: p.name.lower()):
                if not sub.is_dir():
                    continue
                slug = sub.name.lower()
                tech = ORCH_TECH_SUBS.get(slug)
                if tech:
                    trees.append((sub, "ORCH", tech))
                    found_tech = True
            if not found_tech:
                trees.append((child, "ORCH", None))
        elif "philharmonia" in name or "philarmonia" in name:
            trees.append((child, "PHIL", None))
        elif "mcgill" in name:
            trees.append((child, "MCGILL", None))
        elif iowa_name and woodwind:
            trees.append((child, "IOWA", "ordinario"))
    return trees


def find_family_donor_trees(root: Path, instrument: str | None) -> list[tuple[Path, str, str | None]]:
    """Sibling-instrument trees used only to build a family-transfer ORCH spine."""
    trees: list[tuple[Path, str, str | None]] = []
    if not root.exists() or not instrument:
        return trees
    for child in sorted(root.iterdir(), key=lambda p: p.name.lower()):
        if not child.is_dir() or not _folder_is_family_donor(child.name, instrument):
            continue
        name = child.name.lower()
        if "orch" in name or "orchidea" in name:
            found_tech = False
            for sub in sorted(child.iterdir(), key=lambda p: p.name.lower()):
                if not sub.is_dir():
                    continue
                tech = ORCH_TECH_SUBS.get(sub.name.lower())
                if tech:
                    trees.append((sub, "ORCH", tech))
                    found_tech = True
            if not found_tech:
                trees.append((child, "ORCH", "ordinario"))
        elif name.startswith("iowa") or "iowa_" in name or "yowa" in name:
            trees.append((child, "IOWA", "ordinario"))
        elif "philharmonia" in name or "philarmonia" in name:
            trees.append((child, "PHIL", "ordinario"))
        elif "mcgill" in name:
            trees.append((child, "MCGILL", "ordinario"))
    return trees


def _path_is_tasto(parts: list[str]) -> bool:
    """Tasto compiled books are the technique sul tasto, never ordinario."""
    return any("tasto" in p for p in parts)


def _tech_from_parts(parts: list[str]) -> str | None:
    if _path_is_tasto(parts):
        return "sul tasto"
    if any(p in {"con-sord", "con-sordino", "muted"} or "con-sord" in p for p in parts):
        return "con sordino"
    if any(p in {"sul-ponticello", "arco-sul-ponticello"} or "ponticello" in p for p in parts):
        return "sul ponticello"
    if any(p == "harmonics" or p.startswith("harmonics") for p in parts):
        return "harmonics"
    if any(p in {"arco-normal", "ordinario", "non-vibrato", "normal"} for p in parts):
        return "ordinario"
    return None


def collect_compiled_specs(trees: list[tuple[Path, str, str | None]]) -> tuple[list[dict], list[str]]:
    specs: list[dict] = []
    warnings: list[str] = []
    for tree, collection, forced_tech in trees:
        warnings.extend(audit_research_tree(tree, collection=collection, technique=forced_tech or ""))
        if not tree.exists():
            continue
        for path in tree.rglob("*compiled_density_metrics_research.xlsx"):
            if path.name.startswith("~$"):
                continue
            parts = [p.lower() for p in path.parts]
            if _path_is_tasto(parts):
                tech = "sul tasto"
            else:
                tech = forced_tech or _tech_from_parts(parts)
                # Philharmonia / McGill woodwind trees are often dynamic-only folders.
                if not tech and collection in {"PHIL", "MCGILL"}:
                    tech = "ordinario"
            if not tech:
                continue
            dyn = _dyn_from_parts(parts)
            if dyn is None and collection == "MCGILL":
                dyn = "mf"
            if dyn is None:
                continue
            placed = str(path).lower().replace("/", "\\")
            if "_sustains\\" in placed and "_sustains_stable" not in placed:
                continue
            specs.append(
                {
                    "path": path,
                    "collection": collection,
                    "technique": tech,
                    "dynamic": dyn,
                }
            )
    log_discovered_specs(specs)
    return specs, warnings


def discover_in(
    root: Path, *, instrument: str | None = None, media_sheet: str | None = None
) -> tuple[Path | None, list[dict], list[str]]:
    print(f"\nSEARCH  rebuild folder\n  {root}")
    if not root.exists():
        return None, [], [f"missing rebuild folder: {root}"]
    arco = find_zenodo_arco(root, instrument)
    if arco:
        print(f"  FOUND  Zenodo arco  {arco}")
    else:
        print("  FOUND  (no Zenodo collections arco/media workbook)")
    trees = find_research_trees(root, instrument)
    donors = find_family_donor_trees(root, instrument)
    for extra in _sibling_family_deposits(root, instrument):
        for item in find_family_donor_trees(extra, instrument):
            if item[1] in {"IOWA", "ORCH"}:
                donors.append(item)
    if not trees:
        print("  FOUND  (no target Orchidea / Philharmonia / McGill / Iowa trees)")
    for tree, coll, tech in trees:
        print(f"  FOUND  {coll:6} {tech or '(from folder names)':16}  {tree}")
    for tree, coll, tech in donors:
        print(f"  DONOR  {coll:6} {tech or '(from folder names)':16}  {tree}")
    warnings: list[str] = []
    if arco:
        warnings.extend(audit_iowa_zenodo_workbook(arco, media_sheet=media_sheet))
    specs, tree_warnings = collect_compiled_specs(trees)
    donor_specs, donor_warnings = collect_compiled_specs(donors)
    for spec in donor_specs:
        spec["role"] = "family_donor"
    warnings.extend(tree_warnings)
    warnings.extend(donor_warnings)
    discover_in.last_donors = donor_specs  # type: ignore[attr-defined]
    return arco, specs, warnings


def _val(payload) -> float:
    return float(payload[1] if isinstance(payload, tuple) else payload)


def _lab(midi: int, payload=None) -> str:
    if isinstance(payload, tuple):
        return payload[0]
    return midi_to_label(midi)


def _best_measured(
    specs: list[dict],
    *,
    title: str = "FILES USED  (merged compiled books per collection × technique × dynamic)",
    extra_observations=None,
) -> dict[tuple[str, str, str], dict]:
    """Union every matching compiled workbook. Longest-file-wins is prohibited."""
    from ste_lab.empirical import merge_measured_specs

    measured, reports = merge_measured_specs(
        specs, title=title, extra_observations=extra_observations
    )
    _best_measured.last_reports = reports  # type: ignore[attr-defined]
    return measured


def _as_map(curve: dict) -> dict[int, float]:
    return {int(m): _val(p) for m, p in curve.items()}


def _anchors(
    arco: dict[int, float],
    effect: dict,
    *,
    effect_name: str,
    grade: str,
    min_midi: int | None = None,
    dynamic: str = "mf",
    collection: str = "",
) -> pd.DataFrame:
    from ste_lab.relations import collection_from_grade

    rows = []
    coll = collection or collection_from_grade(grade) or ""
    for midi in sorted(set(arco) & set(effect)):
        if min_midi is not None and midi < min_midi:
            continue
        src = float(arco[midi])
        tgt = _val(effect[midi])
        if src <= 0 or tgt <= 0:
            continue
        note = _lab(midi, effect[midi])
        rows.append(
            {
                "effect": effect_name,
                "evidence_grade": grade,
                "collection": coll,
                "dynamic": dynamic,
                "source_cond": "ordinario",
                "target_cond": effect_name,
                "note": note,
                "midi": int(midi),
                "sourceCDM": src,
                "targetCDM": tgt,
                "L": math.log(tgt / src),
                "STE_Lab_paste": f"{note}  {src:.6f}  {tgt:.6f}",
            }
        )
    return pd.DataFrame(rows)


def _get(measured, coll, tech, dyn) -> dict:
    return measured.get((coll, tech, dyn), {}).get("curve", {})


EFFECT_SLUGS = {
    "harmonics": "harmonics",
    "sul ponticello": "ponticello",
    "con sordino": "sordino",
    "sul tasto": "tasto",
}


def effect_column(tech: str, dyn: str) -> str:
    slug = EFFECT_SLUGS.get(tech, tech.replace(" ", "_"))
    return f"ORCH_{slug}_{dyn}"


def orchidea_recordings(measured: dict) -> dict[tuple[str, str], dict]:
    """Every Orchidea effect × dynamic that was actually compiled. No ordinario."""
    out: dict[tuple[str, str], dict] = {}
    for (coll, tech, dyn), blob in measured.items():
        if coll != "ORCH":
            continue
        if tech in {"ordinario", "arco"}:
            continue
        curve = blob.get("curve") or {}
        if curve:
            out[(tech, dyn)] = curve
    return out


def _finite_cell(rec: dict, key: str) -> bool:
    v = rec.get(key, float("nan"))
    return v == v


def _overlap_scale(donor: dict[int, float], target: dict[int, float]) -> float:
    """DEPRECATED leakage scalar. Use calibration.legacy_target_scale."""
    from ste_lab.calibration import legacy_target_scale

    return legacy_target_scale(donor, target)


def _pairing_L_instr(measured: dict, donor_measured: dict, dyn: str, midis: list[int]):
    """Same-collection L_instr = ln(target / sibling). Philharmonia first, then McGill."""
    from ste_lab.calibration import _anchors
    from ste_lab.transfer import _interp_log_ratio

    for coll in ("PHIL", "MCGILL"):
        tgt = _as_map(_get(measured, coll, "ordinario", dyn)) or _as_map(
            _get(measured, coll, "ordinario", "mf")
        )
        sib = _as_map(_get(donor_measured, coll, "ordinario", dyn)) or _as_map(
            _get(donor_measured, coll, "ordinario", "mf")
        )
        anchors = _anchors(sib, tgt)
        if anchors:
            return _interp_log_ratio(anchors, midis), anchors, coll
    return {}, [], None


def write_woodwind_media(
    root: Path,
    instrument: str,
    measured: dict,
    donor_measured: dict,
    media_sheet: str | None,
    sources: set[Path] | None = None,
) -> Path | None:
    """Write IOWA (measured) + ORCH (family-transfer estimate) media book for a woodwind."""
    from ste_lab.calibration import (
        TransferEstimate,
        apply_log_ratio,
        legacy_scalar_transfer,
        load_config,
        remember_transfer_estimates,
        two_log_ratio_transfer,
        _anchors,
    )
    from ste_lab.transfer import _interp_log_ratio

    spec = resolve_instrument(instrument)
    if not spec:
        return None
    cfg = load_config()
    iowa = {d: _as_map(_get(measured, "IOWA", "ordinario", d)) for d in CORE}
    iowa_measured = any(iowa.values())
    orch_donor = {d: _as_map(_get(donor_measured, "ORCH", "ordinario", d)) for d in CORE}
    iowa_sib = {d: _as_map(_get(donor_measured, "IOWA", "ordinario", d)) for d in CORE}
    lo, hi = spec.sounding_low, spec.sounding_high
    orch: dict[str, dict[int, float]] = {d: {} for d in CORE}
    estimates = {}
    iowa_origin = "measured"
    if iowa_measured:
        for dyn in CORE:
            donor = orch_donor[dyn] or orch_donor.get("mf") or {}
            target = iowa[dyn] or iowa.get("mf") or {}
            sibling = iowa_sib[dyn] or iowa_sib.get("mf") or {}
            # Do not clip the transfer grid by catalog sounding_high; Iowa extrema stay.
            midis = sorted(donor)
            if cfg.transfer_method == "legacy_target_calibrated":
                est = legacy_scalar_transfer(donor, target, midis, lo, hi)
                print(
                    f"  legacy_target_calibrated  ORCH {dyn}  scale={est.scale_legacy:.4f}  "
                    f"n={len(est.curve)}  COMBINED_ESTIMATE (deprecated)"
                )
            else:
                est = two_log_ratio_transfer(donor, sibling, target, midis)
                print(
                    f"  two_log_ratio  ORCH {dyn}  L_instr={est.L_instr_status}  "
                    f"anchors_instr={est.n_L_instr_anchors}  n={len(est.curve)}  "
                    f"excluded_from_validation={est.validation_excluded}"
                )
            orch[dyn] = est.curve
            estimates[dyn] = est
    else:
        sib = family_donor_instrument(instrument) or "sibling"
        if not any(orch_donor.values()) and not any(iowa_sib.values()):
            print(
                f"  (no measured Iowa of this instrument and no {sib} Iowa/Orchidea "
                "to apply a Philharmonia/McGill family ratio — cannot write media spine)"
            )
            return None
        for dyn in CORE:
            donor = orch_donor[dyn] or orch_donor.get("mf") or {}
            sibling = iowa_sib[dyn] or iowa_sib.get("mf") or {}
            midis = sorted(set(donor) | set(sibling))
            L_instr, instr_anchors, pair_coll = _pairing_L_instr(
                measured, donor_measured, dyn, midis
            )
            if not L_instr:
                print(
                    f"  two_log_ratio  {dyn}  L_instr=UNDETERMINED  "
                    f"(no {sib} overlap in Philharmonia/McGill)"
                )
                estimates[dyn] = TransferEstimate(
                    curve={},
                    method="two_log_ratio",
                    L_instr_status="UNDETERMINED",
                    notes=(
                        "L_instr UNDETERMINED: no same-collection sibling overlap "
                        "in Philharmonia or McGill. No transfer cells."
                    ),
                )
                continue
            orch[dyn] = apply_log_ratio(donor, L_instr, midis) if donor else {}
            iowa[dyn] = apply_log_ratio(sibling, L_instr, midis) if sibling else {}
            coll_anchors = _anchors(sibling, donor) if sibling and donor else []
            L_coll = _interp_log_ratio(coll_anchors, midis) if coll_anchors else {}
            est = TransferEstimate(
                curve=orch[dyn],
                method="two_log_ratio",
                L_coll=L_coll,
                L_instr=L_instr,
                n_L_coll_anchors=len(coll_anchors),
                n_L_instr_anchors=len(instr_anchors),
                L_instr_status="estimated",
                fitted_on=[f"{pair_coll}_sibling", f"{pair_coll}_target"],
                validation_excluded=[pair_coll],
                notes=(
                    f"L_instr from {pair_coll} target vs {pair_coll} {sib}; "
                    f"IOWA_target = IOWA_{sib} × exp(L_instr); "
                    f"ORCH_target = ORCH_{sib} × exp(L_instr). "
                    f"{pair_coll} is excluded from validation of this transfer."
                ),
            )
            estimates[dyn] = est
            print(
                f"  two_log_ratio  L_instr from {pair_coll} vs {sib}  {dyn}  "
                f"anchors_instr={est.n_L_instr_anchors}  "
                f"n_iowa={len(iowa[dyn])}  n_orch={len(orch[dyn])}  "
                f"excluded_from_validation={est.validation_excluded}"
            )
        if not any(iowa.values()) and not any(orch.values()):
            print("  (Philharmonia/McGill family ratio produced no transfer cells)")
            return None
        iowa_origin = "family_transfer"
    write_woodwind_media.last_estimates = estimates  # type: ignore[attr-defined]
    remember_transfer_estimates(estimates)
    midis = sorted(set().union(*iowa.values(), *orch.values()))
    if not midis:
        return None
    sheet = media_sheet or f"{spec.display_name}_Media"
    gen_dir = Path(root) / "generated"
    gen_dir.mkdir(parents=True, exist_ok=True)
    path = gen_dir / f"{spec.display_name.replace(' ', '_')}_generated_media.xlsx"
    from ste_lab.empirical import assert_not_source

    assert_not_source(path, sources or set())
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme["A1"] = f"{spec.display_name} — woodwind media spine (not a string Zenodo book)"
    origin_name = (
        "combined_estimate"
        if cfg.transfer_method == "legacy_target_calibrated"
        else "orchidea_family_transfer_estimate"
    )
    readme["A3"] = (
        (
            "IOWA columns are measured compiled spectral_mass from this instrument. "
            if iowa_measured
            else "IOWA columns are family transfer: sibling Iowa × exp(L_instr) from a "
            "Philharmonia/McGill same-collection pair. Not an Iowa recording of this instrument. "
        )
        + "ORCH is a deprecated alias for orchidea_family_transfer_estimate: "
        "a model-derived curve, not an Orchidea recording of this instrument. "
        f"Transfer method: {cfg.transfer_method}. "
        "No by-string sheets: woodwinds are not bowed strings."
    )
    readme["A3"].alignment = Alignment(wrap_text=True)
    readme.merge_cells("A3:F6")
    ws = wb.create_sheet(sheet)
    headers = ["Note", "IOWA pp", "IOWA mf", "IOWA ff", "ORCH pp", "ORCH mf", "ORCH ff"]
    ws.append(headers)
    for midi in midis:
        row = [midi_to_label(midi)]
        for coll, curves in (("IOWA", iowa), ("ORCH", orch)):
            for dyn in CORE:
                row.append(curves[dyn].get(midi))
        ws.append(row)
    prov = wb.create_sheet("Spine_provenance")
    prov.append(["collection", "origin", "source"])
    prov.append(
        [
            "IOWA",
            iowa_origin,
            "target Iowa compiled research"
            if iowa_measured
            else "family transfer from sibling Iowa × exp(L_instr)",
        ]
    )
    prov.append(["ORCH", origin_name, f"{cfg.transfer_method}; deprecated column name ORCH"])
    fit = wb.create_sheet("Transfer_fit")
    fit.append(
        [
            "dynamic",
            "midi",
            "transfer_estimate",
            "L_instr",
            "L_coll",
            "method",
            "n_L_coll_anchors",
            "n_L_instr_anchors",
            "L_instr_status",
            "validation_excluded",
            "notes",
        ]
    )
    for dyn, est in estimates.items():
        if not est.curve:
            fit.append(
                [
                    dyn,
                    None,
                    None,
                    None,
                    None,
                    est.method,
                    est.n_L_coll_anchors,
                    est.n_L_instr_anchors,
                    est.L_instr_status,
                    ";".join(est.validation_excluded),
                    est.notes,
                ]
            )
            continue
        for midi, val in sorted(est.curve.items()):
            fit.append(
                [
                    dyn,
                    midi,
                    val,
                    est.L_instr.get(midi),
                    est.L_coll.get(midi),
                    est.method,
                    est.n_L_coll_anchors,
                    est.n_L_instr_anchors,
                    est.L_instr_status,
                    ";".join(est.validation_excluded),
                    est.notes,
                ]
            )
    from ste_lab.relation_export import attach_relation_sheets
    from ste_lab.relations import (
        catalog_for_discovery,
        discover_relations,
        field_for_transfer,
        select_production_relation,
    )
    from ste_lab.validation import compare_relations, summary_validation_lines

    catalog = catalog_for_discovery(
        instrument=instrument,
        donor_measured=donor_measured,
        donor_instrument=family_donor_instrument(instrument),
    )
    ww_relations = discover_relations(measured, catalog)
    select_production_relation(
        ww_relations,
        {
            "instrument": instrument,
            "family": "woodwinds",
            "donor_instrument": family_donor_instrument(instrument),
        },
    )
    if cfg.transfer_field_mode == "pooled":
        for dyn in CORE:
            est = estimates.get(dyn)
            if est is None:
                continue
            midis = sorted(set(est.curve) | set(est.L_instr) | set(est.L_coll))
            instr_peers = [r for r in ww_relations if r.kind == "instrument" and r.dynamic == dyn]
            prod_i = next((r for r in instr_peers if r.used_in_production), None)
            L_instr, contrib_i, msgs_i = field_for_transfer(
                instr_peers,
                midis,
                production=prod_i,
                mode="pooled",
                weight=cfg.transfer_field_weight,
                min_collections=cfg.transfer_field_min_collections,
            )
            coll_peers = [
                r
                for r in ww_relations
                if r.kind == "collection"
                and r.dynamic == dyn
                and r.source_cond == "IOWA"
                and r.target_cond == "ORCH"
            ]
            prod_c = next((r for r in coll_peers if r.used_in_production), None)
            L_coll, _contrib_c, msgs_c = field_for_transfer(
                coll_peers,
                midis,
                production=prod_c,
                mode="pooled",
                weight=cfg.transfer_field_weight,
                min_collections=cfg.transfer_field_min_collections,
            )
            donor = orch_donor[dyn] or orch_donor.get("mf") or {}
            if L_instr:
                orch[dyn] = apply_log_ratio(donor, L_instr, midis) if donor else est.curve
                est.curve = orch[dyn]
                est.L_instr = L_instr
            if L_coll:
                est.L_coll = L_coll
            extra = "; ".join(msgs_i + msgs_c)
            if extra:
                est.notes = (est.notes or "") + " " + extra
            if iowa_origin != "measured" and L_instr:
                sibling = iowa_sib[dyn] or iowa_sib.get("mf") or {}
                iowa[dyn] = apply_log_ratio(sibling, L_instr, midis) if sibling else iowa[dyn]
        write_woodwind_media.last_estimates = estimates  # type: ignore[attr-defined]
        remember_transfer_estimates(estimates)
        midis = sorted(set().union(*iowa.values(), *orch.values()))
        # rewrite the media sheet values after pooling
        if "Note" in [c.value for c in ws[1]]:
            for row in ws.iter_rows(min_row=2):
                ws.delete_rows(2)
            for midi in midis:
                row = [midi_to_label(midi)]
                for coll, curves in (("IOWA", iowa), ("ORCH", orch)):
                    for dyn in CORE:
                        row.append(curves[dyn].get(midi))
                ws.append(row)
    pairwise, spread = compare_relations(ww_relations)
    attach_relation_sheets(
        wb,
        pairwise=pairwise,
        spread=spread,
        summary_lines=summary_validation_lines(pairwise),
    )
    write_woodwind_media.last_relations = ww_relations  # type: ignore[attr-defined]
    write_woodwind_media.last_pairwise = pairwise  # type: ignore[attr-defined]
    write_woodwind_media.last_spread = spread  # type: ignore[attr-defined]
    wb.save(path)
    write_woodwind_media.last_origins = {  # type: ignore[attr-defined]
        "IOWA": iowa_origin,
        "ORCH": origin_name,
    }
    return path


def build(root: Path, instrument: str = "viola", out_path: Path | None = None) -> list[str]:
    """Hunt `root` and write {Instrument}_STE_preparatory.xlsx there."""
    global last_path, last_warnings
    last_path = None
    last_warnings = []
    write_woodwind_media.last_origins = None  # type: ignore[attr-defined]
    write_woodwind_media.last_estimates = None  # type: ignore[attr-defined]
    from ste_lab.calibration import remember_transfer_estimates
    remember_transfer_estimates(None)
    root = Path(root)
    spec = resolve_instrument(instrument)
    instr_id = spec.instrument_id if spec else (instrument or "viola").strip().lower().replace(" ", "_")
    label = display_name(instrument)
    defaults = deposit_defaults(instr_id) or {}
    media_sheet = defaults.get("media_sheet")
    floor = harm_floor(instr_id)
    out = out_path or (root / f"{label.replace(' ', '_')}_STE_preparatory.xlsx")

    print(f"Rebuild preparatory workbook")
    print(f"  instrument  {instr_id} ({label})")
    print(f"  look in     {root}")
    print(f"  write       {out}")
    if floor is not None:
        print(f"  harmonics floor MIDI {floor}")
    root_l = str(root).lower()
    warnings: list[str] = []
    if instr_id == "cello" and "viola" in root_l:
        warnings.append(f"Instrument is cello but the search folder looks like viola: {root}")
    if instr_id == "viola" and "cello" in root_l:
        warnings.append(f"Instrument is viola but the search folder looks like cello: {root}")

    arco_path, specs, hunt_warnings = discover_in(
        root, instrument=instr_id, media_sheet=media_sheet
    )
    warnings.extend(hunt_warnings)
    sources: set[Path] = set()
    if arco_path:
        sources.add(Path(arco_path).resolve())
    for spec in specs:
        sources.add(Path(spec["path"]).resolve())
    extra_obs = []
    if arco_path and orchestral_group(instr_id) == "woodwinds":
        from ste_lab.empirical import iowa_observations_from_media

        extra_obs = iowa_observations_from_media(
            load_arco_media(arco_path, media_sheet=media_sheet),
            Path(arco_path),
        )
    measured = _best_measured(
        specs,
        title="FILES USED  (target instrument — merged compiled books; longest-file-wins is prohibited)",
        extra_observations=extra_obs,
    )
    donor_specs = list(getattr(discover_in, "last_donors", []) or [])
    for spec in donor_specs:
        sources.add(Path(spec["path"]).resolve())
    transfer_donors = list(donor_specs)
    for spec in donor_specs:
        coll = str(spec.get("collection", "")).upper()
        sib = family_donor_instrument(instr_id) or "sibling"
        role = (
            f"L_coll/L_instr {sib}"
            if coll == "IOWA"
            else f"ORCH {sib} donor"
            if coll == "ORCH"
            else f"L_instr {sib}"
            if coll in {"PHIL", "MCGILL"}
            else "unused"
        )
        print(f"  DONOR READ  {coll:6} {spec['path']}  ({role})")
    donor_measured = (
        _best_measured(
            transfer_donors,
            title=(
                "DONOR FILES READ  (sibling "
                f"{family_donor_instrument(instr_id) or 'family'} "
                "Iowa + Orchidea + Philharmonia/McGill; "
                "not written as this instrument)"
            ),
        )
        if transfer_donors
        else {}
    )
    if orchestral_group(instr_id) == "woodwinds":
        rewritten = write_woodwind_media(
            root, instr_id, measured, donor_measured, media_sheet, sources=sources
        )
        if rewritten:
            arco_path = rewritten
            print(f"  WROTE  generated woodwind media spine  {arco_path}")
    if not arco_path:
        raise FileNotFoundError(
            f"No Zenodo collections arco/media workbook in:\n{root}\n"
            "Choose the instrument deposit folder (the one that holds "
            "*Zenodo_collections*Arco*.xlsx or *media*.xlsx). "
            "For woodwinds, Iowa compiled trees plus a family donor "
            "(e.g. Orchidea clarinet) are enough — rebuild will write the media book. "
            "If this instrument has no Iowa, a Philharmonia or McGill pair "
            "(this instrument vs its family donor) plus that donor's Iowa and "
            "Orchidea trees is enough."
        )
    arco = load_arco_media(arco_path, media_sheet=media_sheet)

    inv_rows, long_rows = [], []
    for (coll, tech, dyn), blob in sorted(measured.items()):
        inv_rows.append(
            {
                "collection": coll,
                "technique": tech,
                "dynamic": dyn,
                "n": len(blob["curve"]),
                "source": "; ".join(str(p) for p in (blob.get("paths") or [blob.get("path")])),
            }
        )
        for midi, payload in blob["curve"].items():
            long_rows.append(
                {
                    "collection": coll,
                    "technique": tech,
                    "dynamic": dyn,
                    "note": _lab(midi, payload),
                    "midi": int(midi),
                    "CDM": _val(payload),
                    "source": "; ".join(p.name for p in (blob.get("paths") or [blob.get("path")])),
                }
            )
    inventory = pd.DataFrame(inv_rows)
    measured_long = (
        pd.DataFrame(long_rows).sort_values(["technique", "collection", "dynamic", "midi"])
        if long_rows
        else pd.DataFrame()
    )

    orch_arco = {d: arco["ORCH"][d] for d in CORE}
    iowa_arco = {d: arco["IOWA"][d] for d in CORE}
    woodwind = orchestral_group(instr_id) == "woodwinds"
    spine = "ordinario" if woodwind else "arco"
    recorded = orchidea_recordings(measured)
    if woodwind:
        recorded = {
            key: curve
            for key, curve in recorded.items()
            if key[0].lower() not in {"sul ponticello", "con sordino", "harmonics", "sul tasto"}
        }
    mcgill_sord = {} if woodwind else _get(measured, "MCGILL", "con sordino", "mf")
    mcgill_arco = _as_map(_get(measured, "MCGILL", "ordinario", "mf"))
    phil_harm = {} if woodwind else _get(measured, "PHIL", "harmonics", "mf")
    mcgill_harm = {} if woodwind else _get(measured, "MCGILL", "harmonics", "mf")
    phil_ord = {d: _as_map(_get(measured, "PHIL", "ordinario", d)) for d in (*CORE, "p", "mp", "f")}
    mcgill_ord = {d: _as_map(_get(measured, "MCGILL", "ordinario", d)) for d in (*CORE, "p", "mp", "f")}

    from ste_lab.relations import (
        catalog_for_discovery,
        discover_relations,
        inventory_dataframe,
        relation_anchor_frame,
        select_production_relation,
    )

    catalog = catalog_for_discovery(
        instrument=instr_id,
        arco={"IOWA": iowa_arco, "ORCH": orch_arco},
        donor_measured=donor_measured,
        donor_instrument=family_donor_instrument(instr_id),
        harm_floor=floor,
    )
    relations = discover_relations(measured, catalog)
    selected = select_production_relation(
        relations,
        {
            "instrument": instr_id,
            "family": "woodwinds" if woodwind else orchestral_group(instr_id),
            "donor_instrument": family_donor_instrument(instr_id),
        },
    )
    # Production Anchors_all: ORCH teacher at pp/mf/ff, else Philharmonia, else McGill.
    frames = [
        relation_anchor_frame(
            rel,
            target_payloads=_get(measured, rel.collection, rel.target_cond, rel.dynamic)
            or (_get(measured, "MCGILL", rel.target_cond, rel.dynamic) if rel.collection == "MCGILL" else {}),
        )
        for rel in selected
        if rel.kind == "technique"
    ]
    if not frames:
        for (tech, dyn), curve in sorted(recorded.items()):
            teacher = orch_arco.get(dyn) or {}
            min_m = floor if tech == "harmonics" else None
            frames.append(
                _anchors(teacher, curve, effect_name=tech, grade="Empirical_ORCH", min_midi=min_m, dynamic=dyn)
            )
        if (
            not woodwind
            and not any(tech == "con sordino" for tech, _dyn in recorded)
            and mcgill_sord
            and mcgill_arco
        ):
            frames.append(
                _anchors(mcgill_arco, mcgill_sord, effect_name="con sordino", grade="prediction_McGill", dynamic="mf")
            )
    frames = [df for df in frames if df is not None and not df.empty]
    anchors = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    relations_inventory = inventory_dataframe(relations)

    midi_pool = set(mcgill_sord) | set(mcgill_arco) | set(phil_harm) | set(mcgill_harm)
    for dyn in CORE:
        midi_pool |= set(iowa_arco[dyn]) | set(orch_arco[dyn])
    for curve in list(phil_ord.values()) + list(mcgill_ord.values()):
        midi_pool |= set(curve)
    for curve in recorded.values():
        midi_pool |= set(curve)
    midis = sorted(midi_pool)
    grid_rows = []
    for midi in midis:
        rec = {"note": midi_to_label(midi), "midi": int(midi)}
        for coll, curves in (("IOWA", iowa_arco), ("ORCH", orch_arco)):
            for dyn in CORE:
                rec[f"{coll}_{spine}_{dyn}"] = curves[dyn].get(midi, float("nan"))
        ia, oa = rec[f"IOWA_{spine}_mf"], rec[f"ORCH_{spine}_mf"]
        rec[f"Media_{spine}_mf"] = (ia + oa) / 2.0 if ia == ia and oa == oa else float("nan")
        for (tech, dyn), curve in sorted(recorded.items()):
            col = effect_column(tech, dyn)
            if tech == "harmonics" and floor is not None and midi < floor:
                rec[col] = float("nan")
            else:
                rec[col] = _val(curve[midi]) if midi in curve else float("nan")
            spine_key = f"ORCH_{spine}_{dyn}"
            l_key = f"L_{EFFECT_SLUGS.get(tech, tech.replace(' ', '_'))}_{dyn}"
            if _finite_cell(rec, spine_key) and _finite_cell(rec, col):
                rec[l_key] = math.log(rec[col] / rec[spine_key])
            else:
                rec[l_key] = float("nan")
        rec[f"McGill_{spine}_mf"] = mcgill_arco.get(midi, float("nan"))
        if not woodwind:
            rec["McGill_sordino_mf"] = _val(mcgill_sord[midi]) if midi in mcgill_sord else float("nan")
            rec["PHIL_harmonics_mf"] = _val(phil_harm[midi]) if midi in phil_harm else float("nan")
            rec["McGill_harmonics_mf"] = _val(mcgill_harm[midi]) if midi in mcgill_harm else float("nan")
        for dyn, curve in phil_ord.items():
            rec[f"PHIL_ordinario_{dyn}"] = curve.get(midi, float("nan"))
        for dyn, curve in mcgill_ord.items():
            rec[f"McGill_ordinario_{dyn}"] = curve.get(midi, float("nan"))
        if not woodwind:
            if _finite_cell(rec, f"McGill_{spine}_mf") and _finite_cell(rec, "McGill_sordino_mf"):
                rec["L_sordino_McGill"] = math.log(rec["McGill_sordino_mf"] / rec[f"McGill_{spine}_mf"])
            else:
                rec["L_sordino_McGill"] = float("nan")
        grid_rows.append(rec)
    grid = pd.DataFrame(grid_rows)

    paste_arco_rows = []
    paste_arco_midis = set()
    for dyn in CORE:
        paste_arco_midis |= set(iowa_arco[dyn]) | set(orch_arco[dyn])
    for midi in sorted(paste_arco_midis):
        rec = {"note": midi_to_label(midi), "midi": int(midi)}
        for coll, curves in (("IOWA", iowa_arco), ("ORCH", orch_arco)):
            for dyn in CORE:
                rec[f"{coll}_{dyn}"] = curves[dyn].get(midi, float("nan"))
        paste_arco_rows.append(rec)
    paste_arco_cols = ["note", "midi"] + [
        f"{coll}_{dyn}" for coll in ("IOWA", "ORCH") for dyn in CORE
    ]
    paste_arco = pd.DataFrame(paste_arco_rows, columns=paste_arco_cols)

    paste_midis = set(mcgill_arco)
    if not woodwind:
        paste_midis |= set(mcgill_sord) | set(phil_harm) | set(mcgill_harm)
    for curve in recorded.values():
        paste_midis |= set(curve)
    for curve in list(phil_ord.values()) + list(mcgill_ord.values()):
        paste_midis |= set(curve)
    paste_eff_rows = []
    for midi in sorted(paste_midis):
        rec = {"note": midi_to_label(midi), "midi": int(midi)}
        for (tech, dyn), curve in sorted(recorded.items()):
            col = effect_column(tech, dyn)
            if tech == "harmonics" and floor is not None and midi < floor:
                rec[col] = float("nan")
            else:
                rec[col] = _val(curve[midi]) if midi in curve else float("nan")
        rec["McGill_ordinario_mf"] = mcgill_arco.get(midi, float("nan"))
        if not woodwind:
            rec["McGill_sordino_mf"] = _val(mcgill_sord[midi]) if midi in mcgill_sord else float("nan")
            rec["PHIL_harmonics_mf"] = _val(phil_harm[midi]) if midi in phil_harm else float("nan")
            rec["McGill_harmonics_mf"] = _val(mcgill_harm[midi]) if midi in mcgill_harm else float("nan")
        for dyn, curve in phil_ord.items():
            rec[f"PHIL_ordinario_{dyn}"] = curve.get(midi, float("nan"))
        for dyn, curve in mcgill_ord.items():
            rec[f"McGill_ordinario_{dyn}"] = curve.get(midi, float("nan"))
        paste_eff_rows.append(rec)
    paste_eff_cols = ["note", "midi"]
    for tech, dyn in sorted(recorded):
        paste_eff_cols.append(effect_column(tech, dyn))
    paste_eff_cols.append("McGill_ordinario_mf")
    if not woodwind:
        paste_eff_cols.extend(
            ["McGill_sordino_mf", "PHIL_harmonics_mf", "McGill_harmonics_mf"]
        )
    paste_effects = (
        pd.DataFrame(paste_eff_rows)
        if paste_eff_rows
        else pd.DataFrame(columns=paste_eff_cols)
    )

    tree_names = ", ".join(
        sorted({t[0].parent.name if t[2] else t[0].name for t in find_research_trees(root, instr_id)})
    ) or "(none)"
    recorded_txt = ", ".join(f"{tech} {dyn}" for tech, dyn in sorted(recorded)) or "(none)"
    floor_txt = f"MIDI {floor}" if floor is not None else "no extra floor"
    woodwind = orchestral_group(instr_id) == "woodwinds"
    arco_origins = getattr(write_woodwind_media, "last_origins", None) or {
        "IOWA": "measured",
        "ORCH": "measured" if not woodwind else "family_transfer",
    }
    wb = Workbook()
    ink = Font(name="Calibri", size=11, color="2B2A26")
    title = Font(name="Calibri Light", size=18, color="1F4E79")
    ws = wb.active
    ws.title = "README"
    ws["A1"] = f"{label} — STE preparatory pack (all grounded effects)"
    ws["A1"].font = title
    iowa_is_measured = str(arco_origins.get("IOWA", "measured")).lower() == "measured"
    ws["A3"] = (
        f"Rebuilt from {root}. "
        + (
            "Woodwind run: no arco, no sul ponticello, no con sordino, no invented harmonics. "
            + (
                "IOWA is measured. ORCH ordinario is family transfer from the sibling Orchidea tree."
                if iowa_is_measured
                else "IOWA and ORCH ordinario are family transfer: L_instr from Philharmonia/McGill "
                "(this instrument vs family donor) applied to the donor Iowa and Orchidea trees."
            )
            if woodwind
            else "One file for harmonics, sul ponticello and con sordino when those trees exist. "
            "Inputs only. Extra-collection technique L is used when Orchidea lacks that effect at pp/mf/ff."
        )
    )
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:F5")
    lines = [
        "Sheets",
        "• All_effects_mf — every mf number side by side"
        + ("." if woodwind else ", plus L for each string effect."),
        "• Paste_arco — IOWA/ORCH ordinario pp/mf/ff (Tab 2).",
        "• Paste_effects_mf — "
        + (
            "Philharmonia / McGill ordinario context. No string-effect columns."
            if woodwind
            else "every Orchidea recording (effect × dynamic) as its own column."
        ),
        "• Anchors_all — filter by effect, copy STE_Lab_paste into Tab 3.",
        "• Relations_inventory — every transfer relation found in the trees (used_in_production = today's teacher).",
        "• Measured_inventory / Measured_long — provenance.",
        "",
        "Orchidea rule",
        *(
            [
                (
                    "• IOWA is measured compiled spectral_mass on this instrument."
                    if iowa_is_measured
                    else "• IOWA ordinario is family transfer (sibling Iowa × exp(L_instr)), not an Iowa recording of this instrument."
                ),
                "• ORCH ordinario is family transfer from the sibling Orchidea tree, not a recording of this instrument.",
                "• String techniques (arco, ponticello, sordino, harmonics) are not written.",
                f"• Recorded here: {recorded_txt}.",
            ]
            if woodwind
            else [
                "• If Orchidea recorded that effect at that dynamic, those cells stay measured. No L on that layer.",
                "• L is only for IOWA, or for an Orchidea dynamic that was not recorded.",
                f"• Recorded here: {recorded_txt}. Harmonics floor {floor_txt}.",
                "• Extra collections teach a technique that IOWA/ORCH lack, at pp/mf/ff only. p/mp/f/fff are not production.",
                "• Philharmonia / McGill harmonics sit on the grid as context. They are not Media.",
            ]
        ),
        "",
        "STE Lab: same anchors for IOWA and ORCH. Do not teach L from IOWA.",
        "Hands-off run: python run_ste_from_preparatory.py --prep this file",
        f"Sources: {arco_path.name}; {tree_names}.",
        f"Regenerate: python build_ste_preparatory.py --root \"{root}\" --instrument {instr_id}",
    ]
    for i, line in enumerate(lines, start=7):
        ws[f"A{i}"] = line
        ws[f"A{i}"].font = ink
    ws.column_dimensions["A"].width = 112

    def put(name: str, df: pd.DataFrame, note: str = "") -> None:
        w = wb.create_sheet(name)
        if note:
            w["A1"] = note
            w["A1"].font = Font(name="Calibri", size=10, italic=True, color="5C5A54")
            start = 3
        else:
            start = 1
        if df is None or df.empty:
            if name == "Anchors_all":
                headers = [
                    "effect", "evidence_grade", "collection", "dynamic",
                    "source_cond", "target_cond", "note", "midi",
                    "sourceCDM", "targetCDM", "L", "STE_Lab_paste",
                ]
            elif df is not None and list(df.columns):
                headers = [str(c) for c in df.columns]
            elif name.startswith("Paste_"):
                headers = ["note", "midi"]
            else:
                w.cell(start, 1, "(empty)")
                return
            for c_i, h in enumerate(headers, start=1):
                cell = w.cell(start, c_i, h)
                cell.fill = PatternFill("solid", fgColor="1F4E79")
                cell.font = Font(name="Calibri", size=10, color="FFFFFF", bold=True)
            return
        for r_i, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start):
            for c_i, val in enumerate(row, start=1):
                cell = w.cell(r_i, c_i, val)
                if isinstance(val, float) and val == val:
                    cell.number_format = "0.000000"
        fill = PatternFill("solid", fgColor="1F4E79")
        hf = Font(name="Calibri", size=10, color="FFFFFF", bold=True)
        for cell in w[start]:
            cell.fill = fill
            cell.font = hf
            cell.number_format = "@"
        for col in w.columns:
            letter = col[0].column_letter
            w.column_dimensions[letter].width = min(52, max(12, len(str(col[0].value or "")) + 3))

    put(
        "All_effects_mf",
        grid,
        (
            "All mf curves on one grid. No L_harmonics / L_ponticello / L_sordino — those string techniques are not written for woodwinds."
            if woodwind
            else "All mf curves on one grid. L_harmonics / L_ponticello / L_sordino are Empirical_ORCH when those columns exist."
        ),
    )
    put("Paste_arco", paste_arco, "Tab 2: Note + the IOWA or ORCH column for the dynamic you are loading.")
    put(
        "Paste_effects_mf",
        paste_effects,
        (
            "Tab 2: McGill/Phil ordinario are context. String-effect columns are omitted."
            if woodwind
            else "Tab 2: every Orchidea recording is a column ORCH_{effect}_{dynamic}. Those cells stay measured. McGill/Phil are context unless they are the only sordino teacher."
        ),
    )
    put(
        "Anchors_all",
        anchors,
        "Tab 3: filter effect, copy STE_Lab_paste. Production outputs stay pp/mf/ff; eligible donor dynamics are kept independently.",
    )
    put(
        "Relations_inventory",
        relations_inventory,
        "Every transfer relation the trees can teach. used_in_production marks the teacher chosen by select_production_relation.",
    )
    put("Measured_inventory", inventory, "Merged compiled files per collection × technique × dynamic (union, not longest-file-wins).")
    put("Measured_long", measured_long, f"Every measured spectral_mass cell in the compiled trees under {root}.")
    put(
        "Spine_provenance",
        pd.DataFrame(
            [
                {"collection": "IOWA", "origin": arco_origins.get("IOWA", "measured")},
                {"collection": "ORCH", "origin": arco_origins.get("ORCH", "measured")},
            ]
        ),
        "How Paste_arco was made. family_transfer is not Empirical_ORCH.",
    )

    from ste_lab.empirical import assert_not_source

    assert_not_source(out, sources)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("wrote", out)
    if not anchors.empty:
        keys = [c for c in ("effect", "evidence_grade", "dynamic") if c in anchors.columns]
        print(anchors.groupby(keys).size().to_string())
    if warnings:
        print(f"\n{len(warnings)} research-tree warning(s). Rebuild continued.")
    last_path = out
    last_warnings = warnings
    return warnings


def main(argv: list[str] | None = None) -> list[str]:
    p = argparse.ArgumentParser(description="Rebuild STE preparatory Excel from a deposit folder")
    p.add_argument("--root", type=Path, default=DEPOSIT_DEFAULTS["viola"]["root"])
    p.add_argument("--instrument", default="viola")
    p.add_argument("--out", type=Path, default=None)
    args = p.parse_args(argv)
    return build(args.root, args.instrument, args.out)


if __name__ == "__main__":
    main()
